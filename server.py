"""漫剧工坊：仅监听本机的素材制作服务。"""
from __future__ import annotations
import argparse
import base64
import concurrent.futures
import copy
import csv
import hashlib
import io
import json
import math
import mimetypes
import os
from pathlib import Path
import platform
import re
import secrets
import shutil
import subprocess
import sys
import threading
import time
import urllib.error
import urllib.parse
import urllib.request
import uuid
import webbrowser
import zipfile
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
import ai_providers as ai
import api_connections as connections
import media_download
import promotions
import task_progress as telemetry
import hongguo_source
import guoapp_source
from source_imports import check_storage
import licensing,local_session,engine_manager,runtime_paths

ROOT = Path(__file__).resolve().parent
APP_HOME = Path(sys.executable).resolve().parent if getattr(sys,'frozen',False) else ROOT
sys.path.insert(0, str(ROOT / 'vendor'))
import track_sheets
import vsr_pool
import jianying_projects
import highlight_render
import lama_local
import publish_workflow
DATA = runtime_paths.data_home()
DATA.mkdir(parents=True, exist_ok=True)
PUBLIC = ROOT / 'web'
LOCK = threading.RLock()
JOBS = {}
QUEUE_THREAD = None
QUEUE_WAKE = threading.Event()
QUEUE_STOP = threading.Event()
SOURCE_MANAGER = None
PUBLISH_WORKSPACE = None
LOCAL_JOB_KINDS = {'pipeline','transcode','subtitle','preview','playback','rewrite','cover','coverAI','records','cost','highlight','package','hook','promotion_package','track_sheet','jianying_project'}
SECRETS = {'apiKey': '', 'subtitleKey': ''}
FFMPEG = os.environ.get('FFMPEG_EXE') or shutil.which('ffmpeg')
if not FFMPEG:
    try:
        import imageio_ffmpeg
        FFMPEG = imageio_ffmpeg.get_ffmpeg_exe()
    except ImportError:
        FFMPEG = None
FONT = next((p for p in ['C:/Windows/Fonts/msyh.ttc', 'C:/Windows/Fonts/simhei.ttf', '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'] if Path(p).exists()), '')
VIDEO_EXT = {'.mp4', '.mov', '.mkv', '.avi', '.webm', '.m4v', '.ts'}
DEFAULT = {
    'source': '', 'sourceImportRoot':'', 'output': str(runtime_paths.output_home()),
    'segmentSeconds': 40, 'regroup': True, 'width': 1080, 'height': 1920,
    'parallel': 2, 'gpu': False, 'brightness': 0, 'contrast': 1, 'saturation': 1,
    'speed': 1, 'cropPercent': 0, 'markAI': False, 'markPosition': 'top-left',
    'autoText': False, 'autoCover': False, 'autoAICover':False, 'autoRecords': True, 'autoCost': False,
    'apiBase': 'https://ark.cn-beijing.volces.com/api/v3', 'textModel': 'deepseek-v4-1-flash-260910', 'imageModel': '',
    'textProvider':'ark', 'textProtocol':'chat', 'arkServiceTier':'default', 'textThinking':'disabled',
    'imageProvider':'ark', 'imageApiBase':'https://ark.cn-beijing.volces.com/api/v3',
    'imageProtocol':'ark_images', 'imageSize':'2K', 'aiProfiles':{'text':{},'image':{}},
    'aiConnections':[], 'aiRouting':{c:{'primary':'','backups':[],'autoFallback':False} for c in connections.CHANNELS},
    'aiSelection':{c:'auto' for c in connections.CHANNELS},
    'textPrompt': '依据原始简介重写简介，保留事实，不虚构情节。剧名为4至12个汉字，简介不超过200字。',
    'coverPrompt': '为这部漫剧创作竖版海报，保留参考图片中的人物特征，重新设计背景和构图，不保留原有文字，不添加标题文字。',
    'subtitleMode': 'local', 'subtitleRegion': {'x': 0.05, 'y': 0.78, 'w': 0.9, 'h': 0.14},
    'subtitleEndpoint': '', 'subtitlePoll': '', 'subtitleAuto': False, 'subtitleFirst': 4,
    'subtitleRegions': [], 'pipelineHighlight': False,
    'highlightCount': 1, 'highlightMin': 2, 'highlightMax': 6, 'highlightRandom': True,
    'highlightFixed': 4.2, 'highlightParallel': 2, 'highlightLowResource': False,
    'highlightMusic': '', 'highlightTailEnabled': False, 'highlightTailDirectory': '', 'highlightTailChoice': 'random',
    'endingText': '', 'costTemplate': '', 'costRows': [
        {'name': '演员报酬', 'amount': 0}, {'name': '制作与剪辑', 'amount': 0},
        {'name': 'AI服务', 'amount': 0}, {'name': '音乐与素材授权', 'amount': 0}, {'name': '其他', 'amount': 0}],
    'vsrRoot': '', 'vsrPython': '' if getattr(sys,'frozen',False) else sys.executable, 'vsrMode': 'sttn-auto', 'vsrReuseWorker':True, 'lamaDevice':'auto',
    'sauRoot': str(ROOT/'engines'/'social-auto-upload'), 'sauPython': sys.executable,
    'uploadMode': 'draft', 'uploadSchedule': '',
}

def read_json(path, fallback):
    try:
        return json.loads(path.read_text('utf-8'))
    except (OSError, ValueError):
        return copy.deepcopy(fallback)

STATE = read_json(DATA / 'state.json', {'config': {}, 'projects': [], 'accounts': [], 'logs': []})
saved_vsr_mode=STATE.get('config',{}).get('vsrMode')
if 'highlightFixed' not in STATE.get('config',{}):STATE.setdefault('config',{})['highlightFixed']=STATE.get('config',{}).get('highlightMin',DEFAULT['highlightFixed'])
STATE['config'] = DEFAULT | STATE.get('config', {})
connections.retire_video_config(STATE['config'])
# 同一台电脑上的独立 AI 引擎可供各版本共用，模型和密钥不打入桌面包。
portable_engine=engine_manager.active(DATA)
if portable_engine:
    STATE['config'].update(vsrRoot=portable_engine['root'],vsrPython=portable_engine['python'],subtitleMode='vsr')
    if not saved_vsr_mode:STATE['config']['vsrMode']=portable_engine.get('verifiedMode','sttn-auto')
if not STATE['config']['vsrRoot'] and not getattr(sys,'frozen',False):
    for home in [APP_HOME, *list(APP_HOME.parents)[:3]]:
        install = read_json(home / 'engines' / 'vsr-install.json', {})
        if install.get('root') and install.get('python') and Path(install['root']).is_dir() and Path(install['python']).is_file():
            STATE['config'].update(vsrRoot=install['root'], vsrPython=install['python'], subtitleMode='vsr')
            if not saved_vsr_mode:STATE['config']['vsrMode']=install.get('verifiedMode','sttn-auto')
            break
for key in ['projects', 'accounts', 'logs','promotionTasks','promotionRevenue']:
    STATE.setdefault(key, [])
try:STATE['trackSheetOptions']=track_sheets.options(STATE.get('trackSheetOptions',{}))
except ValueError:STATE['trackSheetOptions']=track_sheets.options({})

# 保存的任务只含非密钥配置；进程、取消信号等运行对象不会写入磁盘。
for item in read_json(DATA / 'jobs.json', []):
    if not isinstance(item,dict) or not item.get('id') or not isinstance(item.get('config'),dict):continue
    if 'highlightFixed' not in item['config']:item['config']['highlightFixed']=item['config'].get('highlightMin',DEFAULT['highlightFixed'])
    item['config'] = connections.retire_video_config(DEFAULT | {k:v for k,v in item['config'].items() if k in DEFAULT})
    if item.get('kind')=='video_generate':
        item['moduleRetired']=True
        if item.get('status') not in ('completed','failed','cancelled'):
            item.update(status='cancelled',stage='视频生成功能已移除，已停止本机查询；如有云端任务请到服务商核对。',finished=time.time())
    if item.get('status') in ['queued','running','pausing']:
        item.update(status='interrupted',stage='上次运行已中断，确认后继续',interruptedAt=time.time())
    item.update(cancel=threading.Event(),processes=set())
    JOBS[item['id']] = item

def replace_state_file(temporary,target):
    # Windows readers and antivirus may briefly hold a file without delete sharing.
    # Keep the previous complete JSON intact and retry only the atomic replace.
    for attempt in range(6):
        try:temporary.replace(target);return
        except PermissionError:
            if attempt==5:raise
            time.sleep(.01*(2**attempt))


def persist_jobs():
    with LOCK:
        rows=[{k:v for k,v in j.items() if k not in ['cancel','processes']} for j in JOBS.values()]
        path=DATA/'jobs.tmp'
        path.write_text(json.dumps(rows,ensure_ascii=False,indent=2),'utf-8')
        replace_state_file(path,DATA/'jobs.json')

def save():
    with LOCK:
        tmp = DATA / 'state.tmp'
        tmp.write_text(json.dumps(STATE, ensure_ascii=False, indent=2), 'utf-8')
        replace_state_file(tmp,DATA / 'state.json')
        persist_jobs()

def log(message, level='info', job=None):
    entry = {'time': time.strftime('%Y-%m-%d %H:%M:%S'), 'message': ai.redact(message), 'level': level, 'job': job}
    with LOCK:
        STATE['logs'].append(entry)
        STATE['logs'] = STATE['logs'][-1000:]
        save()

def number(value, low, high, name):
    try:
        n = float(value)
    except (ValueError, TypeError):
        raise ValueError(f'{name}必须是数字')
    if not math.isfinite(n) or not low <= n <= high:
        raise ValueError(f'{name}须在 {low} 至 {high} 之间')
    return n

def validate_config(c):
    ai.validate(c)
    connections.validate(c)
    if not isinstance(c['aiProfiles'],dict):raise ValueError('AI平台配置记录无效')
    # Only non-secret connection fields may be persisted in provider profiles.
    profile_fields={'textProvider','textProtocol','apiBase','textModel','arkServiceTier','textThinking',
                    'imageProvider','imageApiBase','imageProtocol','imageModel','imageSize'}
    clean_profiles={'text':{},'image':{}}
    for channel in clean_profiles:
        items=c['aiProfiles'].get(channel,{})
        if not isinstance(items,dict):raise ValueError('平台配置记录无效')
        for provider,profile in items.items():
            if provider not in {p['id'] for p in ai.PRESETS} or not isinstance(profile,dict):raise ValueError('平台配置记录无效')
            if any(k not in profile_fields for k in profile):raise ValueError('平台配置只允许保存地址、模型与协议，不能保存密钥')
            clean_profiles[channel][provider]=profile
    c['aiProfiles']=clean_profiles
    if not isinstance(c.get('sourceImportRoot'),str) or len(c['sourceImportRoot'])>32767:raise ValueError('采集保存目录格式无效')
    for key,value in DEFAULT.items():
        if isinstance(value,bool) and not isinstance(c[key],bool):raise ValueError(f'{key}必须为布尔值')
    if c['subtitleMode'] not in ['local','cloud','vsr']:raise ValueError('字幕模式无效')
    if c['vsrMode'] not in ['sttn-auto','lama-box','sttn','sttn-det','lama','propainter']:raise ValueError('本地修复引擎无效')
    if c['lamaDevice'] not in ['auto','cpu']:raise ValueError('LaMa 推理设备无效')
    if c['uploadMode'] not in ['draft','now','scheduled']:raise ValueError('上传模式无效')
    for key, lo, hi in [('segmentSeconds', 1, 3600), ('width', 64, 3840), ('height', 64, 3840),
                        ('parallel', 1, 8), ('brightness', -.2, .2), ('contrast', .8, 1.3),
                        ('saturation', 0, 2), ('speed', .5, 2), ('cropPercent', 0, 15),
                        ('highlightCount', 1, 10), ('highlightMin', 0, 10), ('highlightMax', 0, 10), ('highlightFixed',0,10), ('highlightParallel',1,8),
                        ('subtitleFirst', 1, 100)]:
        number(c[key], lo, hi, key)
    for key in ['width', 'height', 'parallel', 'highlightCount', 'highlightParallel', 'subtitleFirst']:
        if float(c[key]) != int(c[key]):
            raise ValueError(f'{key}必须为整数')
        c[key] = int(c[key])
    if c['width'] % 2 or c['height'] % 2:
        raise ValueError('输出宽高必须为偶数')
    if c['highlightMin'] > c['highlightMax']:
        raise ValueError('高光倍率最小值不能超过最大值')
    highlight_render.validate(c)
    if c['markPosition'] not in ['top-left', 'top-right', 'bottom-left', 'bottom-right']:
        raise ValueError('AI标识位置无效')
    if not isinstance(c['subtitleRegions'],list) or len(c['subtitleRegions'])>8:
        raise ValueError('字幕选区最多8个')
    for r in [c['subtitleRegion'],*c['subtitleRegions']]:
        if not isinstance(r,dict):raise ValueError('字幕区域格式无效')
        for key in ['x', 'y', 'w', 'h']:number(r.get(key), 0, 1, '字幕区域')
        if r['w'] <= 0 or r['h'] <= 0 or r['x'] + r['w'] > 1.000001 or r['y'] + r['h'] > 1.000001:
            raise ValueError('字幕区域超出画面')
    for row in c['costRows']:
        number(row['amount'], 0, 1e12, '成本金额')
    if any(c.get(k) and not c[k].startswith('https://') for k in ['apiBase', 'subtitleEndpoint', 'subtitlePoll']):
        raise ValueError('外部服务地址必须使用 HTTPS，证书校验始终启用')
    out = Path(c['output']).expanduser().resolve()
    if str(out) == out.anchor or out == ROOT or out == DATA:
        raise ValueError('请选择单独的输出文件夹')
    c['output'] = str(out)
    return c

def run_process(args, job=None, timeout=None, cwd=None, observe=None, track=None, env=None):
    if job and JOBS[job]['cancel'].is_set():
        raise InterruptedError('任务已停止')
    proc = subprocess.Popen([str(x) for x in args], stdout=subprocess.PIPE, stderr=subprocess.PIPE, cwd=cwd, env=env,
                            creationflags=getattr(subprocess, 'CREATE_NO_WINDOW', 0))
    if job:
        with LOCK:
            JOBS[job]['processes'].add(proc)
            if JOBS[job]['cancel'].is_set():proc.kill()
    try:
        if observe:
            # Drain both pipes concurrently, including tqdm carriage-return updates. Never block on a full stderr pipe.
            import codecs
            tails={'out':bytearray(),'err':bytearray()};reader_errors=[]
            def drain(pipe,key):
                decoder=codecs.getincrementaldecoder('utf-8')('replace');pending=''
                try:
                    while chunk:=pipe.read1(8192):
                        tails[key].extend(chunk)
                        if len(tails[key])>65536:del tails[key][:-65536]
                        pending+=decoder.decode(chunk)
                        parts=re.split(r'[\r\n]',pending);pending=parts.pop()[-65536:]
                        for line in parts:
                            if line:observe(line)
                    if pending:observe(pending)
                except Exception as e:reader_errors.append(e)
            readers=[threading.Thread(target=drain,args=(proc.stdout,'out'),daemon=True),threading.Thread(target=drain,args=(proc.stderr,'err'),daemon=True)]
            for reader in readers:reader.start()
            begun=time.monotonic()
            while proc.poll() is None:
                if job and JOBS[job]['cancel'].is_set():proc.kill()
                if timeout and time.monotonic()-begun>timeout:raise subprocess.TimeoutExpired(args,timeout)
                if job:telemetry.heartbeat(JOBS[job],track)
                try:proc.wait(timeout=.25)
                except subprocess.TimeoutExpired:pass
            for reader in readers:reader.join(timeout=5)
            if reader_errors:raise reader_errors[0]
            stdout,stderr=bytes(tails['out']),bytes(tails['err'])
        else:
            stdout, stderr = proc.communicate(timeout=timeout)
        if job and JOBS[job]['cancel'].is_set():
            raise InterruptedError('任务已停止')
        if proc.returncode:
            msg = stderr.decode('utf-8', 'replace')[-3000:]
            raise RuntimeError('媒体处理失败：' + msg)
        return stdout
    except subprocess.TimeoutExpired:
        proc.kill()
        proc.communicate()
        raise RuntimeError('媒体命令超时')
    finally:
        if job:
            with LOCK:
                JOBS[job]['processes'].discard(proc)

def ff(args, job=None):
    if not FFMPEG:
        raise RuntimeError('缺少 FFmpeg，请运行安装依赖.ps1')
    measured=job and str(args[-1]).lower().endswith(('.mp4','.mkv','.mov'))
    if not measured:return run_process([FFMPEG,'-hide_banner','-loglevel','error','-y',*args],job)
    target=str(args[-1]);source=next((str(args[i+1]) for i,x in enumerate(args[:-1]) if x=='-i'),'')
    total=None
    try:
        if '-t' in args:total=float(args[args.index('-t')+1])
        elif source.endswith('concat.txt'):total=sum(float(x) for x in re.findall(r'^duration ([0-9.]+)',Path(source).read_text('utf-8'),re.M))
        elif Path(source).is_file():total=probe(source)['duration']
    except (OSError,ValueError,RuntimeError):pass
    video_copy=('-c:v' in args and args[args.index('-c:v')+1]=='copy')
    label=('合并切分（视频流复用）' if video_copy else '重组分集与编码') if '%' in target else '画面与音轨编码'
    tid=telemetry.begin(JOBS[job],label,source,total)
    mode=args[args.index('-c:v')+1] if '-c:v' in args else '未指定'
    with telemetry.LOCK:
        row=telemetry._row(JOBS[job],tid)
        if row:row['encoder']=gpu_encoding.ENCODER_LABELS.get(mode,mode)
    result=run_process([FFMPEG,'-hide_banner','-loglevel','error','-nostdin','-nostats','-stats_period','0.5','-progress','pipe:1','-y',*args],job,
        observe=lambda line:telemetry.ffmpeg_event(JOBS[job],tid,line),track=tid)
    telemetry.finalizing(JOBS[job],tid,'编码已结束，正在校验输出')
    outputs=list(Path(target).parent.glob(re.sub(r'%0?\d*d','*',Path(target).name))) if '%' in target else [Path(target)]
    if not outputs:raise RuntimeError('媒体处理没有生成输出')
    for output in outputs:
        if probe(output)['duration']<=0:raise RuntimeError('输出视频时长无效')
    telemetry.finish(JOBS[job],tid)
    return result

def probe(path):
    if not FFMPEG:
        raise RuntimeError('尚未安装 FFmpeg，无法扫描视频')
    proc = subprocess.run([FFMPEG, '-hide_banner', '-i', str(path)], capture_output=True,
                          creationflags=getattr(subprocess, 'CREATE_NO_WINDOW', 0), timeout=30)
    t = proc.stderr.decode('utf-8', 'replace')
    dur = re.search(r'Duration: (\d+):(\d+):([\d.]+)', t)
    dim = re.search(r'Video:.*?\b(\d{2,5})x(\d{2,5})\b', t)
    if not dur or not dim:
        raise ValueError(f'无法读取视频：{Path(path).name}')
    h, m, s = dur.groups()
    return {'duration': int(h)*3600 + int(m)*60 + float(s), 'width': int(dim[1]),
            'height': int(dim[2]), 'audio': 'Audio:' in t,
            'codec': (re.search(r'Video:\s*(\w+)',t).group(1) if re.search(r'Video:\s*(\w+)',t) else '')}

import gpu_encoding
GPU_SELECTION_LOCK=threading.RLock()

def environment(force=False):
    gpu=gpu_encoding.probe(FFMPEG,force=force)
    gpu['adapters']=gpu_encoding.adapters(force=force)
    return {'ffmpeg': bool(FFMPEG), 'gpu': gpu, 'cpu': platform.processor(), 'cores': os.cpu_count(),
            'python': platform.python_version(), 'font': bool(FONT), 'data': str(DATA)}


def natural(s):
    return [int(x) if x.isdigit() else x.casefold() for x in re.split(r'(\d+)', str(s))]

def get_project(pid):
    p = next((p for p in STATE['projects'] if p['id'] == pid), None)
    if not p:
        raise ValueError('素材项目不存在，请重新扫描')
    return p

def scan(source):
    root = Path(source).expanduser().resolve()
    if not root.is_dir():
        raise ValueError('源文件夹不存在')
    output = Path(STATE['config']['output']).resolve()
    groups = {}
    for path in root.rglob('*'):
        if path.suffix.lower() in VIDEO_EXT and path.is_file() and not path.is_relative_to(output):
            groups.setdefault(path.parent, []).append(path)
    if not groups:
        raise ValueError('该文件夹中未找到支持的视频，请检查输出目录是否包含源目录')
    found, warnings = [], []
    for directory, files in sorted(groups.items(), key=lambda x: natural(x[0])):
        videos = []
        for f in sorted(files, key=natural):
            try:
                videos.append({'path': str(f), 'name': f.name, 'bytes': f.stat().st_size, 'mtime':f.stat().st_mtime, **probe(f)})
            except Exception as e:
                warnings.append(str(e))
        if not videos:
            continue
        pid = hashlib.sha256(str(directory).casefold().encode()).hexdigest()[:16]
        old = next((p for p in STATE['projects'] if p['id'] == pid), {})
        intro = ''
        for f in sorted(directory.glob('*.txt')):
            try:
                intro = f.read_text('utf-8-sig')[:20000]
                break
            except UnicodeError:
                try:
                    intro = f.read_text('gb18030')[:20000]
                    break
                except UnicodeError:
                    continue
        covers = sorted([f for f in directory.iterdir() if f.suffix.lower() in {'.jpg', '.png', '.jpeg', '.webp'}])
        p = {'id': pid, 'source': str(directory), 'name': old.get('name', directory.name),
             'originalName': directory.name, 'intro': old.get('intro', intro), 'originalIntro': intro,
             'cover': str(covers[0]) if covers else '', 'videos': videos, 'status': old.get('status', '未处理'),
             'artifacts': old.get('artifacts', []), 'review': old.get('review', '待准备'),
             'rights': old.get('rights', False), 'notes': old.get('notes', '')}
        found.append(p)
    if not found:
        raise ValueError('找到的视频均无法解析，请检查格式或文件完整性')
    with LOCK:
        ids = {p['id'] for p in found}
        if any(ids.intersection(j['ids']) and j['status'] in ['queued','running','pausing'] for j in JOBS.values()):
            raise ValueError('该素材有正在处理或排队的任务，请先暂停或停止任务再重新扫描')
        for p in found:
            old = next((x for x in STATE['projects'] if x['id']==p['id']), {})
            for key in ['finalVideoPaths','finalJob','sourcePreviewPaths','promotionEntry','promotionEntries','promotionClips','sourceMetadata']:
                if key in old:p[key]=old[key]
        STATE['projects'] = [p for p in STATE['projects'] if p['id'] not in ids] + found
        STATE['config']['source'] = str(root)
        save()
    log(f'扫描完成：{len(found)} 部剧，{sum(len(p["videos"]) for p in found)} 个视频')
    for warning in warnings:
        log(warning, level='warning')
    return {'projects': found, 'warnings': warnings}

def import_track_files(paths):
    if not isinstance(paths,list) or not 1<=len(paths)<=track_sheets.MAX_FILES:raise ValueError('一次请选择1至30个视频文件')
    chosen=[];seen=set()
    for value in paths:
        if not isinstance(value,str) or not Path(value).is_absolute():raise ValueError('请选择有效的本地视频路径')
        path=Path(value).resolve();key=str(path).casefold()
        if key in seen:raise ValueError('所选视频有重复')
        seen.add(key)
        if not path.is_file() or path.suffix.lower() not in VIDEO_EXT:raise ValueError('不是支持的视频文件：'+path.name)
        chosen.append({'path':str(path),'name':path.name,'bytes':path.stat().st_size,'mtime':path.stat().st_mtime,**probe(path)})
    imported=[]
    with LOCK:
        groups={}
        for v in chosen:groups.setdefault(str(Path(v['path']).parent),[]).append(v)
        for directory,videos in groups.items():
            pid=hashlib.sha256(str(directory).casefold().encode()).hexdigest()[:16]
            if any(pid in j['ids'] and j['status'] in ['queued','running','pausing'] for j in JOBS.values()):raise ValueError('该目录有正在处理的任务，请先暂停后导入')
        for directory,videos in groups.items():
            pid=hashlib.sha256(str(directory).casefold().encode()).hexdigest()[:16]
            p=next((p for p in STATE['projects'] if p['id']==pid),None)
            if p is None:
                name=Path(directory).name
                p={'id':pid,'source':directory,'name':name,'originalName':name,'intro':'','originalIntro':'','cover':'','videos':[],
                   'status':'未处理','artifacts':[],'review':'待准备','rights':False,'notes':''}
                STATE['projects'].append(p)
            by_path={v['path']:v for v in p['videos']}
            by_path.update({v['path']:v for v in videos});p['videos']=sorted(by_path.values(),key=lambda x:natural(x['name']))
            imported.extend({'projectId':pid,**v} for v in videos)
        save()
    return {'videos':imported}

def track_request(body):
    opts=track_sheets.options(body.get('trackSheet',STATE.get('trackSheetOptions',{})))
    if opts['output'] and Path(opts['output']).resolve() in {ROOT.resolve(),DATA.resolve(),APP_HOME.resolve()}:raise ValueError('请选择独立的工程图输出目录')
    ids=body.get('ids');requested=body.get('sourceVideoPaths')
    if not isinstance(ids,list) or not ids or len(set(ids))!=len(ids):raise ValueError('请先选择有效的剧集素材')
    if not isinstance(requested,dict) or set(requested)!=set(ids):raise ValueError('所选视频与剧集不匹配')
    fixed={};videos=[]
    for pid in ids:
        p=get_project(pid);available={v['path']:v for v in p['videos']};paths=requested[pid]
        if not isinstance(paths,list) or not paths or any(not isinstance(x,str) or x not in available for x in paths) or len(set(paths))!=len(paths):raise ValueError('请选择属于素材库的有效视频')
        fixed[pid]=list(paths);videos.extend({'projectId':pid,**available[path]} for path in paths)
    if not 1<=len(videos)<=track_sheets.MAX_FILES:raise ValueError('一次最多生成30集工程图')
    order=body.get('trackOrder',[v['path'] for v in videos])
    if not isinstance(order,list) or len(order)!=len(videos) or any(not isinstance(x,str) for x in order) or set(order)!={v['path'] for v in videos}:raise ValueError('分轨处理顺序无效')
    if body.get('sheetMode')=='single' and len(videos)!=1:raise ValueError('单张模式请选择1个视频')
    mapping={v['path']:v for v in videos};videos=[mapping[path] for path in order]
    return {'trackSheet':opts,'sourceVideoPaths':fixed,'trackOrder':order},videos

def track_plan(body):
    request,videos=track_request(body);opts=request['trackSheet'];files=[];errors=[]
    for i,v in enumerate(videos):
        if not Path(v['path']).is_file():errors.append(v['name']+' 已被移动或删除')
        target=track_sheets.output_directory(v,opts)/track_sheets.filename(v,opts,i)
        folder=target.parent
        try:
            existing=folder
            while not existing.exists() and existing!=existing.parent:existing=existing.parent
            if shutil.disk_usage(existing).free<64*1024**2:errors.append(v['name']+' 输出磁盘可用空间不足64MB')
        except OSError:errors.append(v['name']+' 输出目录不可用')
        files.append({'name':v['name'],'outputName':target.name,'directory':str(folder),'duration':v['duration']})
    return {'ok':not errors,'errors':errors,'count':len(videos),'files':files,'tracks':opts['tracks'],'width':opts['width']}

def track_sheet_batch(ids,cfg,jid):
    job=JOBS[jid];opts=job['options']['trackSheet'];order=job['options']['trackOrder']
    videos={v['path']:(p,v) for pid in ids for p in [get_project(pid)] for v in p['videos']}
    total=len(order);done=0;failures=[];results=job.setdefault('trackResults',[])
    telemetry.batch(job,'track_sheets','分轨工程图',total,done)
    for index,path in enumerate(order):
        if job['cancel'].is_set():raise InterruptedError('分轨任务已停止')
        p,video=videos[path];job['project']=p['name'];key=f'{p["id"]}:tracks:'+hashlib.sha256(path.encode()).hexdigest()[:16]
        previous=next((r for r in results if r.get('sourcePath')==path),None);files=cached(jid,key)
        if files and previous:
            try:
                if sha256(previous['path'])!=previous['imageSHA256'] or sha256(previous['manifest'])!=previous['manifestSHA256']:raise ValueError('工程图或分析记录校验不符')
            except (OSError,ValueError,KeyError):files=None
        if files and previous:
            done+=1;telemetry.batch(job,'track_sheets','分轨工程图',total,done);stage(jid,f'复用已完成工程图：{video["name"]}',done/total*100);continue
        progress={'phase':None,'tid':None}
        labels={'hash':('核对源视频','bytes'),'device':('检测解码设备','files'),'decode':('解码与抽帧','seconds'),'analyze':('识别镜头分段','frames'),'cache':('复用视频分析','files'),'render':('绘制并保存PNG','files')}
        def notify(phase,value,length):
            if job['cancel'].is_set():raise InterruptedError('分轨任务已停止')
            if phase=='backend':job['trackDevice']=value;return
            if progress['phase']!=phase:
                if progress['tid']:telemetry.finish(job,progress['tid'])
                label,unit=labels.get(phase,('分析视频','seconds'))
                progress['phase']=phase;progress['tid']=telemetry.begin(job,label,video['name'],length,unit)
                stage(jid,f'{index+1}/{total} · {label} · {video["name"]}',done/total*100)
            telemetry.advance(job,progress['tid'],value,length)
        try:
            result=track_sheets.generate(video,opts,DATA/'track-sheet-cache',FFMPEG,
                lambda args,observe=None,timeout=None:run_process(args,jid,timeout=timeout,observe=observe,track=progress['tid']),
                job['cancel'],index,notify)
            result.update(projectId=p['id'],sourcePath=path)
            register(p,Path(result['path']),'分轨工程图',jid);register(p,Path(result['manifest']),'分轨分析记录',jid)
            results[:]=[r for r in results if r.get('sourcePath')!=path];results.append(result)
            checkpoint(jid,key,[result['path'],result['manifest']]);done+=1
            if progress['tid']:telemetry.finish(job,progress['tid'])
            job['trackDevice']=result['device'];telemetry.batch(job,'track_sheets','分轨工程图',total,done)
            log(('复用缓存，' if result['cacheHit'] else '')+'已生成分轨工程图：'+video['name'],job=jid)
        except InterruptedError:raise
        except Exception as error:
            failures.append({'name':video['name'],'message':str(error)[:400]});log('分轨失败：'+video['name']+' · '+str(error)[:200],level='warning',job=jid)
        job['trackErrors']=list(failures);persist_jobs()
    job['trackErrors']=failures;job['result']={'files':[r['path'] for r in results],'completed':len(results),'failed':len(failures)}
    if failures:raise ValueError(f'已完成{len(results)}张工程图，{len(failures)}集失败；可重试未完成部分。'+failures[0]['name']+'：'+failures[0]['message'])
    if 'track_sheet' not in job['completedStages']:job['completedStages'].append('track_sheet')

def jianying_request(body):
    opts=jianying_projects.options(body.get('jianying',STATE.get('jianyingOptions',{})))
    ids=body.get('ids');kind=body.get('sourceKind','original')
    if kind=='original':
        request,videos=track_request(body)
    elif kind in ('final','subtitle'):
        if not isinstance(ids,list) or not ids or len(set(ids))!=len(ids):raise ValueError('请先选择剧集')
        videos=[{'projectId':pid,**v,**probe(v['path'])} for pid in ids for v in hook_inputs(get_project(pid),kind)]
        if not 1<=len(videos)<=30:raise ValueError('请选择1至30个已生成视频；可以将特定成品导入素材库后自由选集')
        request={}
    else:raise ValueError('工程输入版本无效')
    for video in videos:
        st=Path(video['path']).stat();video['inputStat']={'bytes':st.st_size,'mtime':st.st_mtime_ns}
    return request|{'jianying':opts,'jianyingVideos':videos,'sourceKind':kind},videos


def jianying_batch(ids,cfg,jid):
    job=JOBS[jid];opts=job['options']['jianying'];videos=job['options']['jianyingVideos'];results=job.setdefault('jianyingResults',[])
    root=Path(opts['output'] or str(Path(cfg['output'])/'剪映工程'))
    analysis_opts=track_sheets.options(STATE.get('trackSheetOptions',{}))
    for index,video in enumerate(videos):
        telemetry.batch(job,'jianying_projects','剪映真实工程',len(videos),index)
        if job['cancel'].is_set():raise InterruptedError('工程导出已停止')
        st=Path(video['path']).stat()
        if video['inputStat']!={'bytes':st.st_size,'mtime':st.st_mtime_ns}:raise ValueError('工程输入视频已变化，请重新创建任务')
        key='jianying:'+hashlib.sha256(video['path'].encode()).hexdigest()[:20]
        old=next((r for r in results if r['sourcePath']==video['path']),None)
        if old and cached(jid,key):
            try:jianying_projects.check(old['directory']);continue
            except (OSError,ValueError):pass
        progress={'phase':None,'last':0,'tid':None}
        def notify(phase,value,length):
            if job['cancel'].is_set():raise InterruptedError('工程导出已停止')
            if phase=='backend':return
            now=time.monotonic()
            if progress['phase']==phase and now-progress['last']<.5 and value!=length:return
            if progress['phase']!=phase:
                if progress['tid']:telemetry.finish(job,progress['tid'])
                progress['tid']=telemetry.begin(job,'剪映工程 · '+phase,video['name'],length or 0,'bytes' if phase in ('copy','hash') else 'files' if phase in ('audio','draft','cache') else 'frames')
            progress.update(phase=phase,last=now)
            if isinstance(value,(int,float)) and isinstance(length,(int,float)):telemetry.advance(job,progress['tid'],value,length)
            stage(jid,f'{index+1}/{len(videos)} · 剪映工程 · {video["name"]} · '+{'copy':'复制素材','audio':'提取原声','draft':'保存草稿','decode':'分析镜头','hash':'核对素材','cache':'复用镜头分析'}.get(phase,'分析视频'),index/len(videos)*90)
        run=lambda args,observe=None,timeout=None:run_process(args,jid,observe=observe,timeout=timeout)
        if opts['splitScenes']:
            analysis,_,_=track_sheets.analyze(video,analysis_opts,DATA/'track-sheet-cache',FFMPEG,run,job['cancel'],notify)
        else:analysis={'sourceSHA256':track_sheets.source_hash(Path(video['path']),job['cancel'],notify),'segments':[{'start':0,'end':video['duration']}]}
        result=jianying_projects.generate(video,analysis,root,FFMPEG,run,job['cancel'],notify)
        if progress['tid']:telemetry.finish(job,progress['tid'])
        telemetry.batch(job,'jianying_projects','剪映真实工程',len(videos),index+1)
        result.update(projectId=video['projectId'],sourcePath=video['path']);p=get_project(video['projectId'])
        with LOCK:
            register(p,Path(result['path']),'剪映真实工程',jid);register(p,Path(result['manifest']),'剪映工程清单',jid)
            results[:]=[r for r in results if r['sourcePath']!=video['path']];results.append(result)
            checkpoint(jid,key,[x for x in Path(result['directory']).rglob('*') if x.is_file()]);persist_jobs()
    job['result']={'completed':len(results),'files':[r['path'] for r in results],'editorValidated':False}
    if 'jianying_project' not in job['completedStages']:job['completedStages'].append('jianying_project')


def project_dir(p, cfg):
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', '_', p['name']).strip(' .')[:70] or '未命名'
    out = Path(cfg['output']).resolve() / f'{name}_{p["id"][:6]}'
    source = Path(p['source']).resolve()
    if out == source or source.is_relative_to(out):
        raise ValueError('导出目录不能覆盖源文件夹')
    out.mkdir(parents=True, exist_ok=True)
    return out

def register(p, path, kind, job):
    with LOCK:
        original = get_project(p['id'])
        original['artifacts'] = [a for a in original['artifacts'] if a['path'] != str(path)]
        original['artifacts'].append({'path': str(path), 'name': path.name, 'kind': kind,
                               'time': time.strftime('%Y-%m-%d %H:%M:%S'), 'bytes': path.stat().st_size, 'job': job})
        p['artifacts'] = original['artifacts']
        save()

def subtitle_regions(cfg):
    return cfg.get('subtitleRegions') or [cfg['subtitleRegion']]

def checkpoint(job, key, files=()):
    with LOCK:
        JOBS[job].setdefault('checkpoints',{})[key] = {
            'files':[{'path':str(f),'bytes':Path(f).stat().st_size,'mtime':Path(f).stat().st_mtime_ns} for f in files],
            'time':time.time()}
        persist_jobs()

def cached(job, key):
    entry=JOBS[job].get('checkpoints',{}).get(key)
    if entry is None:return None
    for item in entry['files']:
        path=Path(item['path'])
        if not path.is_file() or path.stat().st_size!=item['bytes'] or path.stat().st_mtime_ns!=item['mtime']:
            return None
    return [Path(item['path']) for item in entry['files']]

def video_entries(files):
    return [{'path':str(f),'name':Path(f).name,'bytes':Path(f).stat().st_size,**probe(f)} for f in files]

def final_videos(p):
    if p.get('finalVideoPaths'):
        files={a['path']:a for a in p['artifacts']}
        result=[files[path] for path in p['finalVideoPaths'] if path in files and Path(path).is_file()]
        if len(result)!=len(p['finalVideoPaths']):raise ValueError('当前成片被移动或删除，请重新制作')
        return result
    files=[a for a in p['artifacts'] if a['kind']=='视频' and Path(a['path']).is_file()]
    return [a for a in files if a['job']==files[-1]['job']] if files else []

def set_final(p, files, job):
    with LOCK:
        original=get_project(p['id'])
        original.update(finalVideoPaths=[str(f) for f in files],finalJob=job)
        p.update(finalVideoPaths=original['finalVideoPaths'],finalJob=job)
        save()

def stage(job, label, progress):
    licensing.require(DATA)
    with LOCK:
        if JOBS[job]['cancel'].is_set():
            raise InterruptedError('任务已停止')
        item=JOBS[job]
        count=1 if item['kind'] in ('track_sheet','jianying_project') else max(1,len(item['ids']))
        overall=(item.get('currentIndex',0)+progress/100)/count*100
        item.update(stage=label, progress=round(overall))
        telemetry.phase(item,label)
    log(label, job=job)

def api_json(url, payload, key='', timeout=120):
    headers = {'Content-Type': 'application/json'}
    if key:
        headers['Authorization'] = 'Bearer ' + key
    req = urllib.request.Request(url, data=json.dumps(payload).encode(), headers=headers)
    try:
        with urllib.request.urlopen(req, timeout=timeout) as r:
            return json.load(r)
    except urllib.error.HTTPError as e:
        # 不记录第三方响应正文，避免服务端回显密钥。
        raise RuntimeError(f'外部接口返回 HTTP {e.code}，请检查权限、模型ID和请求配置')
    except urllib.error.URLError:
        raise RuntimeError('外部接口连接失败，请检查网络和 HTTPS 证书')

def rewrite(p, cfg, job):
    if not p['intro'].strip():
        raise ValueError('缺少原始简介，请先填写作品简介')
    system=cfg['textPrompt']+'\n任务是改写而不是提取：必须创作与原剧名不同的新剧名，简介须重新组织表达，保留原有事实，不添加人物、情节或结局。只返回JSON对象，字段 title 和 synopsis，不要代码块。'
    for attempt in range(2):
        response=ai.text_generate(cfg,system, f'原剧名：{p["name"]}\n原简介：{p["intro"]}',max_tokens=1024)
        try:
            text=response['text']
            result=json.loads(re.sub(r'^```(?:json)?\s*|\s*```$', '', text.strip()))
            title,intro=result['title'],result['synopsis']
            if not isinstance(title,str) or not re.fullmatch(r'[\u4e00-\u9fff]{4,12}',title):raise ValueError()
            if not isinstance(intro,str) or not intro.strip() or len(intro)>1000:raise ValueError()
        except (KeyError,IndexError,TypeError,ValueError):
            if attempt==0:
                system+='\n上次返回格式不符，请严格返回合法JSON，新剧名只能包含4至12个汉字。'
                continue
            raise ValueError('模型返回格式不符：剧名需4至12个汉字，简介需为非空文本，请调整提示词后重试')
        if title.strip()!=p['name'].strip() and intro.strip()!=p['intro'].strip():break
        if attempt==1:raise ValueError('模型连续返回原文，未进行实际改写。原信息已保留，请调整提示词或模型后重试。')
        system+='\n上次返回原文，未完成任务。本次必须同时更换剧名和简介表达，禁止照抄输入。'
    with LOCK:
        p['name'], p['intro'] = title, intro
        save()
    out = project_dir(p, cfg) / '剧名简介.json'
    out.write_text(json.dumps({'title': title, 'synopsis': intro, 'model': response['model'], 'provider':response['provider'],
        'seconds':response['seconds'],'usage':response['usage'],'job': job}, ensure_ascii=False, indent=2), 'utf-8')
    register(p, out, '文案', job)


def trace_connection(jid,event):
    with LOCK:
        JOBS[jid].setdefault('aiAttempts',[]).append(event)
        JOBS[jid]['aiAttempts']=JOBS[jid]['aiAttempts'][-100:];persist_jobs()
    if event['event'] in ['fallback','failed','submitted','succeeded','resume']:
        log(f'接口「{event["name"]}」{event["event"]}：{event.get("detail","")}',job=jid,level='warning' if event['event'] in ['failed','fallback'] else 'info')


def get_frame(video, path, at=0, job=None):
    ff(['-ss', str(max(0, at)), '-i', str(video), '-frames:v', '1', '-update', '1', str(path)], job)

def font(size):
    from PIL import ImageFont
    return ImageFont.truetype(FONT, size) if FONT else ImageFont.load_default()

def cover(p, cfg, job, ai=False):
    from PIL import Image, ImageDraw, ImageOps, ImageFilter
    out = project_dir(p, cfg)
    frame = out / '封面底图.jpg'
    if p['cover'] and Path(p['cover']).is_file():
        base = Image.open(p['cover']).convert('RGB')
    else:
        get_frame(p['videos'][0]['path'], frame, min(1, p['videos'][0]['duration']/3), job)
        base = Image.open(frame).convert('RGB')
    if ai:
        import ai_providers
        raw=ai_providers.image_generate(cfg,cfg['coverPrompt']+'\n作品简介：'+p['intro'][:1000],base)
        try:
            base = Image.open(io.BytesIO(raw)).convert('RGB')
        except (KeyError, ValueError, IndexError):
            raise ValueError('图片接口未返回可读取的图片')
        ai_path = out / 'AI封面原图.jpg'
        base.save(ai_path, quality=94)
        register(p, ai_path, 'AI封面', job)
    base = ImageOps.fit(base, (816, 1086), method=Image.Resampling.LANCZOS)
    overlay = Image.new('RGBA', base.size)
    d = ImageDraw.Draw(overlay)
    for y in range(1086):
        alpha = int(230 * max(0, (y - 420) / 666))
        d.line((0,y,816,y), fill=(13,17,36,alpha))
    base = Image.alpha_composite(base.convert('RGBA'), overlay)
    d = ImageDraw.Draw(base)
    title = p['name'][:24]
    lines = [title[i:i+6] for i in range(0,len(title),6)]
    y = 990 - len(lines)*100
    for line in lines:
        d.text((58,y),line,font=font(78),fill='white',stroke_width=1)
        y += 100
    d.text((62,1020),'漫剧 · 剧集封面',font=font(25),fill=(220,225,240))
    path = out / '封面_816x1086.jpg'
    base.convert('RGB').save(path, quality=95)
    register(p,path,'封面',job)

def filters(cfg, info, jobdir):
    # 在缩放前统一帧时间，保留 HEVC 尾帧；缩放可能清除输入帧的 duration。
    vf = [f'setpts=PTS/{cfg["speed"]}', 'fps=30']
    crop = float(cfg['cropPercent'])/100
    if crop:
        vf.append(f'crop=trunc(iw*{1-crop*2}/2)*2:trunc(ih*{1-crop*2}/2)*2')
    # 尺寸一致时跳过无效缩放，保留部分 HEVC 素材的尾帧时长元数据。
    if crop or info.get('width') != cfg['width'] or info.get('height') != cfg['height']:
        vf.extend([f'scale={cfg["width"]}:{cfg["height"]}:force_original_aspect_ratio=decrease',
                   f'pad={cfg["width"]}:{cfg["height"]}:(ow-iw)/2:(oh-ih)/2:color=black'])
    vf.extend(['setsar=1',
               f'eq=brightness={cfg["brightness"]}:contrast={cfg["contrast"]}:saturation={cfg["saturation"]}'])
    if cfg['markAI']:
        from PIL import Image, ImageDraw
        im = Image.new('RGBA',(250,64),(20,22,34,170))
        ImageDraw.Draw(im).text((14,9),'AI创作',font=font(32),fill='white')
        # 路径转义以适配 FFmpeg filtergraph 的盘符与单引号。
        txt = str(Path(FONT)).replace('\\','/').replace(':','\\:').replace("'", "\\'")
        pos = cfg['markPosition']
        x = '24' if pos.endswith('left') else 'w-tw-24'
        y = '24' if pos.startswith('top') else 'h-th-24'
        vf.append(f"drawtext=fontfile='{txt}':text='AI创作':fontsize={max(18,int(cfg['width']/28))}:fontcolor=white:box=1:boxcolor=black@0.5:boxborderw=8:x={x}:y={y}:enable='lt(t,3)'")
    return ','.join(vf)

def encoder_fallback(job,info,runtime=False):
    with LOCK:
        item=JOBS[job];first=not item.get('gpuFallback')
        item['gpuFallback']={'code':info['code'],'detail':info['detail'],'runtime':runtime,'time':time.time()}
        item['encoderSelection']={'args':gpu_encoding.CPU_ARGS.copy(),'name':'libx264'}
        item.setdefault('encoding',{}).update(videoEncoder='libx264',gpuRequested=True,fallbackCode=info['code'],fallbackReason=info['detail'])
        persist_jobs()
    if first:log('GPU 编码不可用，自动使用 CPU：'+info['detail'],level='warning',job=job)


def encoder(cfg,job=None):
    if not cfg['gpu']:return gpu_encoding.CPU_ARGS.copy()
    with GPU_SELECTION_LOCK:
        item=JOBS.get(job,{})
        if item.get('encoderSelection'):return item['encoderSelection']['args'].copy()
        result=gpu_encoding.probe(FFMPEG,cfg['width'],cfg['height'],force=True,
            runner=lambda args:run_process(args,job,timeout=8))
        if result['available']:
            args=result['encoderArgs'].copy()
            if job:
                item['encoderSelection']={'args':args,'name':args[1]}
                item.setdefault('encoding',{}).update(videoEncoder=args[1],gpuRequested=True,probeProfile=result['profile'])
            return args
        if job:encoder_fallback(job,result)
        return gpu_encoding.CPU_ARGS.copy()


def encode_media(args,job):
    args=gpu_encoding.prepare_arguments(args)
    mode=args[args.index('-c:v')+1] if '-c:v' in args else 'unknown'
    if mode in gpu_encoding.HARDWARE_CODECS and JOBS[job].get('gpuFallback'):
        args=gpu_encoding.cpu_arguments(args);mode='libx264'
    try:return ff(args,job),mode
    except RuntimeError as e:
        if mode not in gpu_encoding.HARDWARE_CODECS or not gpu_encoding.startup_failure(e):raise
        if JOBS[job]['cancel'].is_set():raise InterruptedError('任务已停止') from None
        encoder_fallback(job,gpu_encoding.failure(e),runtime=True)
        return ff(gpu_encoding.cpu_arguments(args),job),'libx264'


def encoding_threads(cfg):
    return min(8,max(1,(os.cpu_count() or 2)//max(1,int(cfg['parallel']))))

def transcode(p, cfg, job, subtitle=False):
    complete=cached(job,f'{p["id"]}:video')
    if complete is not None and not subtitle:return complete
    out = project_dir(p,cfg)
    work = out / '.work' / job
    work.mkdir(parents=True,exist_ok=True)
    dest = out / ('字幕处理' if subtitle else '分集') / job
    dest.mkdir(parents=True,exist_ok=True)
    outputs = []
    videos = p['videos'][:int(cfg['subtitleFirst'])] if subtitle else p['videos']
    durations=[round(float(v['duration'])/float(cfg['speed'])*30)/30 for v in videos]
    interval=max(1,round(float(cfg['segmentSeconds'])*30));frames=[round(d*30) for d in durations]
    boundaries=list(range(interval,sum(frames),interval))
    modern=JOBS[job].get('options',{}).get('encodingPlanVersion',1)>=2
    fast=modern and cfg['regroup'] and not subtitle and len(boundaries)<=2000
    compatible=[True];threads=encoding_threads(cfg) if modern else 2
    JOBS[job].setdefault('encoding',{}).update(threadsPerVideo=threads,gpuRequested=bool(cfg['gpu']),
                           regroupMode='video_copy' if fast else 'reencode',planVersion=2 if modern else 1)
    selected_encoder=encoder(cfg,job)
    JOBS[job]['encoding']['videoEncoder']=selected_encoder[1]
    used_encoders=set()
    def encode_one(item):
        i,v = item
        key=f'{p["id"]}:encode-v2:{i}' if fast else f'{p["id"]}:encode:{i}'
        previous=cached(job,key)
        if previous is not None and (not fast or JOBS[job]['checkpoints'][key].get('frameCount')==frames[i]):
            with LOCK:used_encoders.add(JOBS[job]['checkpoints'][key].get('videoEncoder','unknown'))
            return i,previous[0]
        target = work / f'{i:04d}.mp4'
        vf = filters(cfg,v,work)
        if subtitle:
            r=cfg['subtitleRegion']
            # delogo需要在视频内部留出边界，使用比例区域并钳制到有效边界。
            w,h=cfg['width'],cfg['height']
            x=max(1,min(w-3,int(r['x']*w))); y=max(1,min(h-3,int(r['y']*h)))
            rw=max(1,min(w-x-2,int(r['w']*w))); rh=max(1,min(h-y-2,int(r['h']*h)))
            vf += f',delogo=x={x}:y={y}:w={rw}:h={rh}'
        args=(['-threads',str(threads),'-filter_threads','2'] if modern else [])+['-i',v['path']]
        if not v['audio']:
            args += ['-f','lavfi','-i','anullsrc=channel_layout=stereo:sample_rate=48000']
        # 原速时不经过时伸缩滤镜，避免 atempo=1 对声音做不必要的重叠处理。
        af = ('aresample=48000,apad' if float(cfg['speed']) == 1 else
              f'atempo={cfg["speed"]},aresample=48000,apad')
        args += ['-map','0:v:0','-map','0:a:0' if v['audio'] else '1:a:0','-vf',vf,
                 # 以素材时长限制输出并补齐短音轨，避免 -shortest 截去录屏末尾画面。
                 '-af',af,*encoder(cfg,job),'-pix_fmt','yuv420p',
                 '-c:a','aac','-ac','2','-ar','48000','-b:a','128k',
                 '-t',f'{float(v["duration"])/float(cfg["speed"]):.6f}','-threads',str(threads),str(target)]
        if fast:
            offset=sum(frames[:i]);first=(-offset)%interval
            # Container duration is rounded by FFmpeg's metadata display. Align
            # to the selected CFR frame plan, holding at most a short video tail
            # when the audio/container ends a fraction of one frame later.
            args[args.index('-vf')+1]+=',tpad=stop_mode=clone:stop_duration=0.1'
            args[args.index('-t')+1]=f'{durations[i]:.9f}'
            args[-1:-1]=['-frames:v',str(frames[i]),'-force_key_frames',f'expr:gte(n,{first}+n_forced*{interval})','-bf','0','-video_track_timescale','90000']
        progress,used_encoder=encode_media(args,job)
        checkpoint(job,key,[target])
        with LOCK:
            used_encoders.add(used_encoder)
            JOBS[job]['checkpoints'][key]['videoEncoder']=used_encoder
            if fast:
                reported=re.findall(rb'^frame=(\d+)',progress,re.M)
                count=int(reported[-1]) if reported else None
                JOBS[job]['checkpoints'][key]['frameCount']=count
                if count!=frames[i]:compatible[0]=False
                persist_jobs()
            JOBS[job]['telemetry']['batch']['done']+=1
        return i,target
    prefix='encode-v2' if fast else 'encode'
    telemetry.batch(JOBS[job],p['id']+':encode','本阶段已编码的源视频',len(videos),sum(cached(job,f'{p["id"]}:{prefix}:{i}') is not None for i in range(len(videos))))
    with concurrent.futures.ThreadPoolExecutor(max_workers=int(cfg['parallel'])) as executor:
        futures=[executor.submit(encode_one,item) for item in enumerate(videos)]
        for done,future in enumerate(concurrent.futures.as_completed(futures),1):
            outputs.append(future.result())
            stage(job,f'{p["name"]}：已编码 {done}/{len(videos)} 个源视频',15+done/len(videos)*45)
    outputs.sort()
    # CPU and NVENC bitstreams may have different codec parameters.
    if cfg['gpu'] and (len(used_encoders)>1 or 'unknown' in used_encoders):compatible[0]=False
    if cfg['regroup'] and not subtitle:
        # 用相对数字文件名避免源文件名中引号、换行破坏concat清单。
        listing=work/'concat.txt'
        listing.write_text('\n'.join(f"file '{path.name}'\nduration {durations[i]:.9f}" for i,path in outputs),'utf-8')
        if fast and compatible[0]:
            # Global frame-boundary keyframes were placed in the first encode. Only audio is re-encoded.
            ff(['-f','concat','-safe','0','-i',str(listing),'-map','0:v:0','-map','0:a:0','-c:v','copy','-c:a','aac','-threads',str(threads),
                '-af','aresample=48000:async=1:first_pts=0','-t',f'{sum(durations):.9f}',
                '-f','segment','-segment_frames',','.join(map(str,boundaries)) or str(interval),'-reset_timestamps','1','-segment_start_number','1',
                str(dest/'第%03d集.mp4')],job)
        else:
            if fast:
                JOBS[job]['encoding']['regroupMode']='reencode_fallback'
                log('帧数或源片编码器与流复用条件不符，使用兼容重组编码，已保留原片',level='warning',job=job)
            encode_media(['-f','concat','-safe','0','-i',str(listing),*encoder(cfg,job),'-c:a','aac','-threads',str(threads),
                '-af','aresample=48000:async=1:first_pts=0','-t',f'{sum(durations):.9f}',
                '-force_key_frames',f'expr:gte(t,n_forced*{cfg["segmentSeconds"]})','-bf','0',
                '-f','segment','-segment_time',str(cfg['segmentSeconds']),'-segment_time_delta','0.016667','-reset_timestamps','1','-segment_start_number','1',
                str(dest/'第%03d集.mp4')],job)
    else:
        for i,path in outputs:
            shutil.copy2(path,dest/f'第{i+1:03d}集.mp4')
    files=sorted(dest.glob('*.mp4'))
    if not files:
        raise RuntimeError('没有生成分集文件')
    for path in files:
        register(p,path,'字幕处理' if subtitle else '视频',job)
    checkpoint(job,f'{p["id"]}:video',files)
    shutil.rmtree(work)
    log(f'{p["name"]}：输出 {len(files)} 集',job=job)
    return files

def legacy_cloud_subtitle(p,cfg,job):
    endpoint=cfg['subtitleEndpoint']
    if not endpoint or not SECRETS['subtitleKey']:
        raise ValueError('云端AI去字幕尚未配置服务适配器地址及密钥，请查看字幕页的接口契约')
    out=project_dir(p,cfg)/'AI字幕处理'/job
    out.mkdir(parents=True,exist_ok=True)
    if len(subtitle_regions(cfg))>1:raise ValueError('当前云端适配器只支持单字幕选区，请只保留一个选区')
    outputs=[]
    for i,v in enumerate(p['videos'][:cfg['subtitleFirst']]):
        key=f'{p["id"]}:subtitle:{i}'
        previous=cached(job,key)
        if previous is not None:outputs.append(previous[0]);continue
        if Path(v['path']).stat().st_size>80*1024*1024:
            raise ValueError('云端适配器单文件上限80MB，大文件请由服务商提供分片上传适配器')
        result=api_json(endpoint,{'video_base64':base64.b64encode(Path(v['path']).read_bytes()).decode(),
            'filename':v['name'],'region':cfg['subtitleRegion']},SECRETS['subtitleKey'],timeout=300)
        for _ in range(180):
            if JOBS[job]['cancel'].is_set():
                raise InterruptedError('任务已停止')
            if result.get('status') in ['completed','succeeded'] or result.get('video_url'):
                break
            if result.get('status')=='failed':
                raise RuntimeError('云端字幕服务处理失败，请在服务商控制台检查任务')
            if not result.get('task_id') or not cfg['subtitlePoll']:
                raise ValueError('适配器未返回 video_url，异步任务需配置轮询地址')
            JOBS[job]['cancel'].wait(3)
            result=api_json(cfg['subtitlePoll'],{'task_id':result['task_id']},SECRETS['subtitleKey'])
        url=result.get('video_url','')
        if not url.startswith('https://'):
            raise ValueError('字幕适配器未返回HTTPS视频下载地址')
        path=out/f'第{i+1:03d}集.mp4'
        with urllib.request.urlopen(url,timeout=300) as response, path.open('wb') as f:
            while chunk:=response.read(1024*1024):
                if JOBS[job]['cancel'].is_set():
                    raise InterruptedError('任务已停止')
                f.write(chunk)
        probe(path)
        register(p,path,'AI字幕处理',job)
        checkpoint(job,key,[path]);outputs.append(path)
        stage(job,f'AI去字幕：{i+1} 集完成',20+(i+1)/min(len(p['videos']),cfg['subtitleFirst'])*70)
    return outputs


def cloud_subtitle(p,cfg,job):
    if not connections.candidates(cfg,'subtitle'):return legacy_cloud_subtitle(p,cfg,job)
    if len(subtitle_regions(cfg))>1:raise ValueError('云端字幕适配器当前只支持单个选区')
    out=project_dir(p,cfg)/'AI字幕处理'/job;out.mkdir(parents=True,exist_ok=True);outputs=[]
    for i,v in enumerate(p['videos'][:cfg['subtitleFirst']]):
        key=f'{p["id"]}:subtitle:{i}';previous=cached(job,key)
        if previous is not None:outputs.append(previous[0]);continue
        token=f'{p["id"]}:{i}';remote=JOBS[job].get('remoteSubtitles',{}).get(token)
        def record(value):
            with LOCK:JOBS[job].setdefault('remoteSubtitles',{})[token]=value;persist_jobs()
        if remote and remote.get('status') in ['submitting','submission_unknown']:
            raise ai.ProviderError('字幕接口上次提交结果未知，请到服务商核对；未重复上传视频','SubmissionUnknown')
        if not remote or remote.get('status')=='rejected':
            if Path(v['path']).stat().st_size>80*1024*1024:raise ValueError('字幕适配器单文件上限80MB')
            payload={'filename':v['name'],'video_base64':base64.b64encode(Path(v['path']).read_bytes()).decode(),'region':subtitle_regions(cfg)[0]}
            def submit(selected):
                c=selected['_aiConnection'];record({'connection':c,'status':'submitting'})
                try:result=ai.request_json(c['base']+c['submitPath'],payload,connections.require_key(cfg,c),timeout=300)
                except ai.ProviderError as error:
                    uncertain=error.code in ['NetworkError','InvalidResponse'] or (error.status is not None and error.status>=500)
                    record({'connection':c,'status':'submission_unknown' if uncertain else 'rejected'})
                    if uncertain:raise ai.ProviderError('字幕提交结果未知，请到服务商核对；未自动重复提交','SubmissionUnknown') from None
                    raise
                if not result.get('video_url') and not result.get('task_id'):
                    record({'connection':c,'status':'submission_unknown'});raise ai.ProviderError('字幕接口未返回任务ID或结果，请核对服务契约','SubmissionUnknown')
                value={'connection':c,'status':'submitted','result':result};record(value);return value
            remote=connections.routed(cfg,'subtitle',submit)
        c=remote['connection'];result=remote['result']
        for _ in range(600):
            if JOBS[job]['cancel'].is_set():raise InterruptedError('已暂停字幕任务查询')
            if result.get('video_url'):break
            if result.get('status') in ['failed','cancelled']:raise ai.ProviderError('云端字幕任务失败，请核对平台任务','RemoteTaskFailed')
            JOBS[job]['cancel'].wait(3)
            result=ai.request_json(c['base']+c['pollPath'],{'task_id':remote['result']['task_id']},connections.require_key(cfg,c))
        url=result.get('video_url','');target=out/f'第{i+1:03d}集.mp4'
        media_download.download(url,target,JOBS[job]['cancel']);probe(target);register(p,target,'AI字幕处理',job);checkpoint(job,key,[target]);outputs.append(target)
        connections.report(cfg,c,'succeeded');stage(job,f'云端字幕完成 {i+1}/{len(p["videos"])}',15+(i+1)/len(p['videos'])*70)
    return outputs

def local_subtitle(p,cfg,job):
    out=project_dir(p,cfg)/'字幕处理'/job;out.mkdir(parents=True,exist_ok=True)
    outputs=[]
    for i,v in enumerate(p['videos'][:cfg['subtitleFirst']]):
        key=f'{p["id"]}:subtitle:{i}'
        previous=cached(job,key)
        if previous is not None:outputs.append(previous[0]);continue
        target=out/f'第{i+1:03d}集.mp4';vf=[]
        for r in subtitle_regions(cfg):
            w,h=v['width'],v['height']
            x=max(1,min(w-3,int(r['x']*w)));y=max(1,min(h-3,int(r['y']*h)))
            rw=max(1,min(w-x-2,int(r['w']*w)));rh=max(1,min(h-y-2,int(r['h']*h)))
            vf.append(f'delogo=x={x}:y={y}:w={rw}:h={rh}')
        stage(job,f'区域修补：{i+1}/{min(len(p["videos"]),cfg["subtitleFirst"])}',10+i/len(p['videos'])*70)
        ff(['-i',v['path'],'-map','0:v:0','-map','0:a:0?','-vf',','.join(vf),
            '-c:v','libx264','-preset','fast','-crf','18','-c:a','aac','-b:a','192k',
            '-pix_fmt','yuv420p','-threads','2','-movflags','+faststart',str(target)],job)
        probe(target);register(p,target,'字幕处理',job);checkpoint(job,key,[target]);outputs.append(target)
    return outputs

def repair_subtitles(p,cfg,job,all_videos=False):
    options=cfg|{'subtitleFirst':len(p['videos'])} if all_videos else cfg
    return {'local':local_subtitle,'cloud':cloud_subtitle,'vsr':vsr_subtitle}[cfg['subtitleMode']](p,options,job)

def vsr_subtitle(p,cfg,job):
    root=Path(cfg['vsrRoot']).expanduser().resolve()
    runtime=Path(cfg['vsrPython']).expanduser()
    if not (root/'backend'/'main.py').is_file():
        raise ValueError('请先配置 video-subtitle-remover 源码目录，需包含 backend/main.py、模型权重和依赖环境')
    if not runtime.is_file():raise ValueError('字幕引擎Python不存在，请选择已安装该引擎依赖的Python解释器')
    if cfg['vsrMode'] not in ['sttn-auto','lama-box','sttn','sttn-det','lama','propainter']:
        raise ValueError('不支持的字幕修补模型')
    if cfg['vsrMode']=='lama-box' and not lama_local.ready(root,runtime):raise ValueError('LaMa 模型未安装：需要 backend/models/big-lama/big-lama.pt')
    (vsr_pool.POOL if cfg['vsrMode']=='lama-box' else lama_local.POOL).close()
    out=project_dir(p,cfg)/'本地AI去字幕'/job;out.mkdir(parents=True,exist_ok=True)
    videos=p['videos'][:cfg['subtitleFirst']]
    telemetry.batch(JOBS[job],p['id']+':subtitle','本阶段已修复的源视频',len(videos),sum(cached(job,f'{p["id"]}:subtitle:{i}') is not None for i in range(len(videos))))
    outputs=[]
    for i,v in enumerate(videos):
        key=f'{p["id"]}:subtitle:{i}'
        previous=cached(job,key)
        if previous is not None:outputs.append(previous[0]);continue
        w,h=v['width'],v['height']
        target=out/f'第{i+1:03d}集.mp4'
        args=[str(runtime),'-m','backend.main','--input',v['path'],'--output',str(target)]
        for r in subtitle_regions(cfg):
            args+=['--subtitle-area-coords',str(int(r['y']*h)),str(int((r['y']+r['h'])*h)),
                   str(int(r['x']*w)),str(int((r['x']+r['w'])*w))]
        args+=['--inpaint-mode','sttn-det' if cfg['vsrMode']=='sttn' else cfg['vsrMode']]
        stage(job,f'本地AI去字幕：{i+1}/{len(videos)}，模型 {cfg["vsrMode"]}',10+i/len(videos)*80)
        tid=telemetry.begin(JOBS[job],'加载模型／准备字幕修复',v['name'],unit='frames')
        if cfg['vsrMode']=='lama-box' or cfg.get('vsrReuseWorker',True) and cfg['vsrMode']=='sttn-auto':
            coords=[[int(r['y']*h),int((r['y']+r['h'])*h),int(r['x']*w),int((r['x']+r['w'])*w)] for r in subtitle_regions(cfg)]
            def attach(proc):
                with LOCK:JOBS[job]['processes'].add(proc)
            def detach(proc):
                with LOCK:JOBS[job]['processes'].discard(proc)
            pool=lama_local.POOL if cfg['vsrMode']=='lama-box' else vsr_pool.POOL
            stats=pool.run(runtime,root,{'input':v['path'],'output':str(target),'mode':cfg['vsrMode'],'regions':coords,'duration':v['duration'],'device':cfg.get('lamaDevice','auto'),'ffmpeg':FFMPEG},JOBS[job]['cancel'],
                observe=lambda line:telemetry.vsr_event(JOBS[job],tid,line),heartbeat=lambda:telemetry.heartbeat(JOBS[job],tid),attach=attach,detach=detach)
            if not cfg.get('vsrReuseWorker',True):pool.close()
            JOBS[job]['vsrWorker']=stats
            JOBS[job]['localRepairEngine']=cfg['vsrMode']
            if stats.get('fallback'):log('LaMa 推理已自动改用 CPU：'+stats['fallback'],level='warning',job=job)
            log(('复用常驻模型' if stats['reusedProcess'] else '启动常驻模型')+f'，累计处理 {stats["completedVideos"]} 集，模型加载 {stats["modelLoads"]} 次',job=job)
        else:
            vsr_pool.POOL.close()
            run_process(args,job,cwd=root,observe=lambda line:telemetry.vsr_event(JOBS[job],tid,line),track=tid,env=os.environ|{'PYTHONUNBUFFERED':'1','PYTHONIOENCODING':'utf-8'})
        telemetry.finalizing(JOBS[job],tid,'模型处理已退出，正在校验视频')
        if not target.is_file() or target.stat().st_size<128:
            raise RuntimeError('AI引擎未生成有效视频：'+v['name']+'。请检查本地引擎、模型和磁盘空间；本集未完成，已完成集数保留。')
        try:meta=probe(target)
        except (ValueError,RuntimeError) as e:
            raise RuntimeError('AI输出视频无法读取：'+v['name']+'。请查看任务日志中的引擎错误；本集未完成，可修复环境后重试。') from e
        if abs(meta['duration']-v['duration'])>.3:raise RuntimeError('AI输出未覆盖完整视频：'+v['name']+f'，原片 {v["duration"]:.2f} 秒，输出 {meta["duration"]:.2f} 秒。已保留现场，请重试当前视频。')
        if v.get('audio') and not meta['audio']:
            stage(job,'恢复原片音轨：'+v['name'],10+i/len(videos)*80)
            restored=target.with_name(target.stem+'-audio.mp4')
            ff(['-i',str(target),'-i',v['path'],'-map','0:v:0','-map','1:a:0','-c:v','copy','-c:a','aac','-b:a','192k','-movflags','+faststart',str(restored)],job)
            audio_meta=probe(restored)
            if not audio_meta['audio'] or abs(audio_meta['duration']-v['duration'])>.3:
                raise RuntimeError('AI输出音轨恢复失败，本集未登记完成，请查看日志后重试')
            restored.replace(target)
        register(p,target,'AI字幕处理',job);checkpoint(job,key,[target]);outputs.append(target)
        telemetry.finish(JOBS[job],tid)
        with LOCK:JOBS[job]['telemetry']['batch']['done']+=1
    return outputs

def preview_subtitle(p,cfg,job):
    options=JOBS[job]['options'];v=p['videos'][options['videoIndex']]
    out=project_dir(p,cfg)/'短片试处理'/job;out.mkdir(parents=True,exist_ok=True)
    source=out/'处理前.mp4';key=f'{p["id"]}:preview-source'
    if cached(job,key) is None:
        ff(['-ss',str(options['start']),'-i',v['path'],'-t',str(options['seconds']),
            '-map','0:v:0','-map','0:a:0?','-c:v','libx264','-preset','fast','-crf','18','-c:a','aac',
            '-pix_fmt','yuv420p','-threads','2','-movflags','+faststart',str(source)],job)
        register(p,source,'试处理原片',job);checkpoint(job,key,[source])
    sample={**p,'videos':video_entries([source])}
    result=repair_subtitles(sample,cfg,job,True)[0]
    JOBS[job]['result']={'before':str(source),'after':str(result),'videoIndex':options['videoIndex'],
                        'start':options['start'],'seconds':probe(source)['duration'],
                        'mode':cfg['subtitleMode'],'engine':cfg['vsrMode'] if cfg['subtitleMode']=='vsr' else cfg['subtitleMode'],'regions':subtitle_regions(cfg)}
    persist_jobs()

def sha256(path):
    h=hashlib.sha256()
    with Path(path).open('rb') as f:
        while chunk:=f.read(1024*1024):
            h.update(chunk)
    return h.hexdigest()

def records(p,cfg,job):
    from PIL import Image,ImageDraw,ImageOps
    out=project_dir(p,cfg)/'工程记录'/job
    out.mkdir(parents=True,exist_ok=True)
    sheet=Image.new('RGB',(1600,1250),'#f5f6fa'); d=ImageDraw.Draw(sheet)
    d.text((50,30),p['name']+' 制作记录',font=font(40),fill='#24243d')
    d.text((50,96),'源文件抽帧 · 实际处理参数 · SHA256校验',font=font(24),fill='#55556c')
    evidence=[]
    total=sum(v['duration'] for v in p['videos'])
    for i in range(4):
        global_t=total*(i+.5)/4
        rem=global_t
        v=p['videos'][-1]
        for current in p['videos']:
            v=current
            if rem<current['duration']:
                break
            rem-=current['duration']
        at=min(max(0,rem),max(0,v['duration']-.05))
        path=out/f'源帧_{i+1}.jpg'
        get_frame(v['path'],path,at,job)
        evidence.append({'frame':path.name,'source':v['path'],'timestamp':round(at,3),'sourceSHA256':sha256(v['path'])})
        image=ImageOps.pad(Image.open(path).convert('RGB'),(730,420),color='#e8e9ef')
        x=50+(i%2)*770; y=160+(i//2)*500
        sheet.paste(image,(x,y));d.text((x,y+432),f'{v["name"][:28]}  |  {at:.2f}s',font=font(22),fill='#35354e')
        register(p,path,'源帧',job)
    d.text((50,1180),'记录时间 '+time.strftime('%Y-%m-%d %H:%M:%S')+'   仅记录实际操作，不代表平台审核结论',font=font(22),fill='#55556c')
    image=out/'工程记录.jpg';sheet.save(image,quality=94);register(p,image,'工程记录',job)
    manifest=out/'工程记录.json'
    publiccfg={k:v for k,v in cfg.items() if k in DEFAULT}
    manifest.write_text(json.dumps({'project':p['name'],'job':job,'created':time.strftime('%Y-%m-%d %H:%M:%S'),
        'sourceFrames':evidence,'parameters':publiccfg,
        'outputs':[{'file':a['path'],'sha256':sha256(a['path'])} for a in final_videos(p)]},ensure_ascii=False,indent=2),'utf-8')
    register(p,manifest,'工程记录',job)

def costs(p,cfg,job):
    from docx import Document
    from docx.shared import Inches,Pt,RGBColor
    from docx.oxml.ns import qn
    rows=cfg['costRows'];total=sum(float(r['amount']) for r in rows)
    if total<=0:
        raise ValueError('成本合计为0，请填写实际成本后生成')
    out=project_dir(p,cfg)
    csvpath=out/'成本比例表.csv'
    with csvpath.open('w',encoding='utf-8-sig',newline='') as f:
        writer=csv.writer(f);writer.writerow(['项目','金额（元）','占比'])
        for r in rows:
            writer.writerow([r['name'],f'{float(r["amount"]):.2f}',f'{float(r["amount"])/total:.2%}'])
        writer.writerow(['合计',f'{total:.2f}','100.00%'])
    register(p,csvpath,'成本报告',job)
    template=cfg.get('costTemplate','').strip()
    if template and (not Path(template).is_file() or Path(template).suffix.lower()!='.docx'):
        raise ValueError('Word模板不存在或不是.docx文件')
    doc=Document(template) if template else Document()
    style=doc.styles['Normal'];style.font.name='Microsoft YaHei';style.font.size=Pt(11)
    style.element.rPr.rFonts.set(qn('w:eastAsia'),'微软雅黑')
    if template:
        values={'{{剧名}}':p['name'],'{{总成本}}':f'{total:.2f}','{{日期}}':time.strftime('%Y-%m-%d')}
        paragraphs=list(doc.paragraphs)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    paragraphs+=list(cell.paragraphs)
        for para in paragraphs:
            if any(k in para.text for k in values):
                text=para.text
                for key,value in values.items():text=text.replace(key,value)
                if para.runs:
                    para.runs[0].text=text
                    for run in para.runs[1:]:run.text=''
                else:para.add_run(text)
        doc.add_page_break()
    title=doc.add_paragraph(p['name']+' 制作成本报告','Title')
    for run in title.runs:run.font.color.rgb=RGBColor(0,0,0)
    doc.add_paragraph(f'本报告记录本剧填报的制作成本及各项占比。成本合计为 {total:,.2f} 元，金额由填报人核实后用于内部自查或报备。')
    table=doc.add_table(rows=1,cols=3);table.style='Light Shading Accent 1'
    for cell,text in zip(table.rows[0].cells,['成本项目','金额（元）','占比']):cell.text=text
    for r in rows:
        cells=table.add_row().cells
        for cell,text in zip(cells,[r['name'],f'{float(r["amount"]):,.2f}',f'{float(r["amount"])/total:.2%}']):cell.text=text
    cells=table.add_row().cells
    for cell,text in zip(cells,['合计',f'{total:,.2f}','100.00%']):cell.text=text
    doc.add_paragraph('填报日期 '+time.strftime('%Y-%m-%d'))
    doc.add_paragraph('填报人 ____________________    复核人 ____________________')
    path=out/'成本报告.docx';doc.save(path);register(p,path,'成本报告',job)

def highlight(p,cfg,job):
    return highlight_render.render(sys.modules[__name__],p,cfg,job)


def promotion_hook(p,cfg,jid):
    if not p['rights']:raise ValueError('请先确认素材制作与推广授权')
    options=JOBS[jid]['options'];v=options.get('hookSource') or p['videos'][options['videoIndex']]
    stat=Path(v['path']).stat()
    if options.get('hookSourceStat')!={'bytes':stat.st_size,'mtime':stat.st_mtime_ns}:raise ValueError('钩子输入视频已变更，请重新选择后创建任务')
    out=project_dir(p,cfg)/'推广钩子'/jid;out.mkdir(parents=True,exist_ok=True);target=out/'钩子短视频.mp4'
    key=f'{p["id"]}:promotion-hook'
    if cached(jid,key) is None:
        stage(jid,'按选定的剧情时间段制作钩子短视频',20)
        vf='scale=trunc(iw/2)*2:trunc(ih/2)*2,setsar=1'
        if options['cta'] and FONT:
            textfile=out/'引导文字.txt';textfile.write_text(options['cta'],'utf-8')
            fp=FONT.replace('\\','/').replace(':','\\:');tp=str(textfile).replace('\\','/').replace(':','\\:').replace("'","\\'")
            vf+=f",drawtext=fontfile='{fp}':textfile='{tp}':fontsize={max(16,int(v['width']/30))}:fontcolor=white:box=1:boxcolor=black@0.6:x=(w-tw)/2:y=h*0.82:enable='gte(t,{max(0,options['seconds']-2)})'"
        ff(['-ss',str(options['start']),'-i',v['path'],'-t',str(options['seconds']),'-map','0:v:0','-map','0:a:0?',
            '-vf',vf,'-c:v','libx264','-preset','fast','-crf','20','-c:a','aac','-pix_fmt','yuv420p','-threads','2','-movflags','+faststart',str(target)],jid)
        meta=probe(target)
        if abs(meta['duration']-options['seconds'])>.3:raise ValueError('钩子输出时长异常，请重新选择时间段')
        register(p,target,'推广钩子',jid);checkpoint(jid,key,[target])
    else:meta=probe(target)
    with LOCK:
        clips=p.setdefault('promotionClips',[])
        if not any(c['id']==jid for c in clips):
            stat=target.stat();clips.append({'id':jid,'path':str(target),'title':options['title'],'sourcePath':v['path'],'start':options['start'],
                'sourceKind':options.get('hookSourceKind','original'),'seconds':meta['duration'],'cta':options['cta'],'bytes':stat.st_size,'mtimeNs':stat.st_mtime_ns,'sha256':sha256(target),'createdAt':time.time()})
        JOBS[jid]['result']={'clipId':jid,'path':str(target),'seconds':meta['duration']};save()

def promotion_package(p,cfg,jid):
    if not p['rights']:raise ValueError('请先确认素材制作与推广授权')
    snapshot=JOBS[jid]['options']['promotionTask'];clip=snapshot['clip'];path=Path(clip['path'])
    if not path.is_file() or path.stat().st_size!=clip['bytes'] or sha256(path)!=clip['sha256']:raise ValueError('推广素材已修改，请重新制作并建立任务')
    out=project_dir(p,cfg)/'推广挂载'/snapshot['id'];out.mkdir(parents=True,exist_ok=True);target=out/f'推广任务_{jid}.zip'
    manifest={'format':1,'taskId':snapshot['id'],'project':snapshot['projectName'],'account':snapshot['accountName'],
        'entry':snapshot['entry'],'clip':{k:clip[k] for k in ['title','start','seconds','sha256']},
        'videoFile':'钩子短视频.mp4','mountStatus':'pending','platformResult':None,'preparedAt':time.time()}
    with zipfile.ZipFile(target,'w',zipfile.ZIP_DEFLATED,compresslevel=1) as z:
        z.write(path,'钩子短视频.mp4');z.writestr('挂载任务.json',json.dumps(manifest,ensure_ascii=False,indent=2))
        z.writestr('发布核对说明.txt','1. 在指定账号的发布后台上传本包视频。\n2. 选择本任务指定的剧集或小程序入口，核对剧名与平台ID。\n3. 在官方页面确认后由本人发布。\n4. 查看实际作品，核对全集入口可点击且指向正确剧集。\n5. 回到软件填写作品ID或链接及核对说明。\n本包只准备素材与挂载信息，没有执行平台发布、广告播放或收益结算。')
    register(p,target,'推广发布包',jid)
    with LOCK:
        row=promotions.task(STATE,snapshot['id']);row['packagePath']=str(target);row['updatedAt']=time.time();save()
    JOBS[jid]['result']={'taskId':snapshot['id'],'path':str(target)}

def package(p,cfg,job):
    if not p['rights']:
        raise ValueError('请先确认本剧拥有制作与发布授权')
    final=final_videos(p)
    if not final:
        raise ValueError('尚未生成分集视频，请先运行制作流水线')
    if not p['intro'].strip():raise ValueError('请先填写剧集简介')
    artifacts=[]
    for kind in ['封面','工程记录','成本报告']:
        items=[a for a in p['artifacts'] if a['kind']==kind and Path(a['path']).is_file()]
        if items:artifacts.extend(a for a in items if a['job']==items[-1]['job'])
    if not any(a['kind']=='封面' for a in artifacts):raise ValueError('请先制作封面')
    artifacts=final+artifacts
    path=project_dir(p,cfg)/f'发布材料_{job}.zip'
    manifest={'title':p['name'],'synopsis':p['intro'],'rightsConfirmed':p['rights'],
              'status':'prepared','files':[], 'created':time.strftime('%Y-%m-%d %H:%M:%S')}
    with zipfile.ZipFile(path,'w',compression=zipfile.ZIP_DEFLATED,compresslevel=1) as z:
        for i,a in enumerate(artifacts):
            name=f'{a["kind"]}/{i+1:03d}_{Path(a["path"]).name}'
            z.write(a['path'],name);manifest['files'].append({'name':name,'sha256':sha256(a['path'])})
        z.writestr('剧集信息.json',json.dumps(manifest,ensure_ascii=False,indent=2))
        z.writestr('提交说明.txt','此包为发布准备材料。请在视频号官方工作台核对剧集信息后手动上传、提交。软件未执行平台提审，审核状态请按实际结果回填。')
    register(p,path,'发布包',job);p['review']='待提交';save()

def upload_engine(cfg):
    root=Path(cfg['sauRoot']).expanduser().resolve()
    runtime=Path(cfg['sauPython']).expanduser()
    if not (root/'uploader/tencent_uploader/main.py').is_file():raise ValueError('未找到视频号上传模块，请检查开源引擎目录')
    if not runtime.is_file():raise ValueError('上传引擎Python不存在')
    return root,runtime

def bridge_command(runtime):
    if getattr(sys,'frozen',False) and Path(runtime).resolve()==Path(sys.executable).resolve():
        return [runtime,'--engine-worker']
    return [runtime,ROOT/'engine_bridge.py']

def engine_login(cfg,jid):
    root,runtime=upload_engine(cfg)
    account_id=JOBS[jid]['accountIds'][0]
    cookie=DATA/'upload-cookies'/f'{account_id}.json';cookie.parent.mkdir(parents=True,exist_ok=True)
    stage(jid,'请在弹出的官方窗口中本人扫码登录',10)
    run_process([*bridge_command(runtime),'login','--engine',root,'--account',cookie],jid,cwd=root)
    if not cookie.is_file():raise RuntimeError('登录引擎未生成账号会话文件')
    log('上传引擎登录完成，账号会话仅保存在本机',job=jid)

def upload_videos(p,cfg,jid):
    if not p['rights']:raise ValueError('请先确认该剧拥有制作与发布授权')
    root,runtime=upload_engine(cfg)
    files=final_videos(p)
    if not files:raise ValueError('请先生成分集视频')
    accounts=JOBS[jid]['accountIds'];total=len(accounts)*len(files);done=0
    out=project_dir(p,cfg)/'上传回执'/jid;out.mkdir(parents=True,exist_ok=True)
    for account_id in accounts:
        account=next((a for a in STATE['accounts'] if a['id']==account_id),None)
        if not account:raise ValueError('选中的账号已被移除')
        cookie=DATA/'upload-cookies'/f'{account_id}.json'
        if not cookie.is_file():raise ValueError(f'{account["name"]} 尚未通过上传引擎扫码登录')
        for i,a in enumerate(files):
            mode=cfg['uploadMode']
            request={'file':a['path'],'title':f'{p["name"]}第{i+1}集','description':p['intro'],
                     'mode':mode,'schedule':cfg['uploadSchedule'] if mode=='scheduled' else '', 'aiContent':cfg['markAI']}
            request_file=out/f'{account_id}_{i+1:03d}_任务.json'
            request_file.write_text(json.dumps(request,ensure_ascii=False,indent=2),'utf-8')
            stage(jid,f'{account["name"]}：上传 {i+1}/{len(files)} 集，模式 {mode}',5+done/total*90)
            run_process([*bridge_command(runtime),'upload','--engine',root,'--account',cookie,'--request',request_file],jid,cwd=root)
            receipt=out/f'{account_id}_{i+1:03d}_结果.json'
            receipt.write_text(json.dumps({'account':account['name'],'file':a['name'],'mode':mode,
                'result':'uploader_completed','time':time.strftime('%Y-%m-%d %H:%M:%S'),
                'note':'上传引擎完成普通视频流程，不代表短剧机构整剧审核通过'},ensure_ascii=False,indent=2),'utf-8')
            register(p,receipt,'上传回执',jid);done+=1

def playback_proxy(p,cfg,jid):
    index=JOBS[jid]['options']['videoIndex'];video=p['videos'][index];source=Path(video['path']);stat=source.stat()
    signature=hashlib.sha256(f'{source}:{stat.st_size}:{stat.st_mtime_ns}'.encode()).hexdigest()[:20]
    target=DATA/'playback'/f'{signature}.mp4';target.parent.mkdir(parents=True,exist_ok=True)
    if not target.is_file():
        stage(jid,'准备可播放的兼容预览副本',20)
        temp=target.with_suffix('.working.mp4')
        ff(['-i',str(source),'-map','0:v:0','-map','0:a:0?', '-vf',"scale=w='min(1280,iw)':h='min(720,ih)':force_original_aspect_ratio=decrease:force_divisible_by=2",
            '-c:v','libx264','-preset','veryfast','-crf','23','-pix_fmt','yuv420p','-c:a','aac','-threads','2','-movflags','+faststart',str(temp)],jid)
        temp.replace(target)
    register(p,target,'播放预览',jid)
    p.setdefault('sourcePreviewPaths',{})[str(source)]={'path':str(target),'bytes':stat.st_size,'mtime':stat.st_mtime}
    checkpoint(jid,f'{p["id"]}:playback:{index}',[str(target)])
    save()


def production_plan(kind,cfg,continuing=False):
    steps=[]
    if kind=='pipeline':
        if cfg['autoText']:steps.append(['AI 剧名与简介','使用本次任务选择的文字接口',['text']])
        if cfg['subtitleAuto'] and not continuing:steps.append(['先修复字幕','全部所选源视频',['subtitle']])
        steps.append(['复用已完成分集' if continuing else '画面处理与分集','保留已有分集，不重复编码' if continuing else (f"每集 {cfg['segmentSeconds']} 秒" if cfg['regroup'] else '保持源分集'),['video']])
        if cfg['pipelineHighlight']:steps.append(['成片包装','使用本次处理后的分集',['highlight']])
        if cfg['autoCover']:steps.append(['制作封面','生成剧集封面',['cover']])
        if cfg['autoRecords']:steps.append(['工程记录','保存真实帧与文件校验值',['records']])
        if cfg['autoCost']:steps.append(['成本报告','按本次配置生成',['cost']])
    else:
        names={'transcode':('仅合并切分','单步完成后结束',['video']),'subtitle':('仅去字幕','单步完成后结束',['subtitle']),
               'highlight':('高光包装','单独处理当前成片',['highlight']),'cover':('制作封面','单独生成封面',['cover']),
               'coverAI':('AI封面','使用所选生图接口',['cover']),'records':('工程记录','单独生成记录',['records']),
               'cost':('成本报告','单独生成报告',['cost']),'rewrite':('AI文案','单独改写剧名与简介',['text']),
               'package':('整理发布材料','导出本地材料',['package']),
               'jianying_project':('剪映真实工程','复制素材并生成可编辑时间线',['jianying_project'])}
        if kind in names:steps.append(list(names[kind]))
    return steps

def continued_video_files(options,pid):
    entries=options.get('continuedVideoFiles',{}).get(pid,[])
    if not entries:raise ValueError('没有可接续的已完成分集')
    valid={a['path'] for a in get_project(pid)['artifacts'] if a.get('job')==options['continueFromJob'] and a['kind']=='视频'}
    files=[]
    for entry in entries:
        path=Path(entry['path'])
        if str(path) not in valid or not path.is_file():raise ValueError('已完成分集被移动或删除，请核对后重新制作')
        stat=path.stat()
        if stat.st_size!=entry['bytes'] or stat.st_mtime_ns!=entry['mtime']:raise ValueError('已完成分集发生变化，接续任务已停止')
        files.append(path)
    return files

def worker(jid,ids,kind,cfg):
    job=JOBS[jid]
    cfg=copy.deepcopy(cfg);cfg['_aiTraceId']=jid;connections.TRACE_CALLBACKS[jid]=lambda event:trace_connection(jid,event)
    run_start=time.time()
    try:
        licensing.require(DATA)
        job.update(status='running',runStarted=run_start)
        job.pop('encoderSelection',None);job.pop('gpuFallback',None)
        for key in ('fallbackCode','fallbackReason','probeProfile'):job.get('encoding',{}).pop(key,None)
        telemetry.settle(job,'interrupted')
        job['telemetry']={'rows':[]};telemetry.phase(job,'正在核对任务与检查点')
        job.setdefault('started',run_start);persist_jobs()
        if job['cancel'].is_set():raise InterruptedError('任务已停止')
        if job.get('sourceSnapshot')!=source_snapshot(ids,job.get('options',{}).get('sourceVideoPaths')):raise ValueError('源素材已修改或移动，请核对后创建新任务；本任务检查点保留')
        if kind=='login':
            engine_login(cfg,jid)
        if kind=='track_sheet':track_sheet_batch(ids,cfg,jid)
        if kind=='jianying_project':jianying_batch(ids,cfg,jid)
        for index,pid in enumerate(ids if kind not in ('track_sheet','jianying_project') else []):
            job['currentIndex']=index
            base_project=get_project(pid)
            paths=job.get('options',{}).get('sourceVideoPaths',{}).get(pid)
            p=SelectedProject(base_project,paths) if paths else base_project
            job['project']=p['name'];stage(jid,f'开始：{p["name"]}',5)
            if cached(jid,f'{pid}:project-done') is not None:continue
            job['currentCompletedStages']=[]
            def completed(name):
                if name not in job['completedStages']:job['completedStages'].append(name)
                if name not in job['currentCompletedStages']:job['currentCompletedStages'].append(name)
            def step(name,action):
                key=f'{pid}:step:{name}'
                if cached(jid,key) is None:
                    before={a['path'] for a in p['artifacts']}
                    action()
                    files=[a['path'] for a in p['artifacts'] if a['job']==jid and a['path'] not in before and Path(a['path']).is_file()]
                    checkpoint(jid,key,files)
                completed(name)
            if kind in ['pipeline','rewrite'] and (kind=='rewrite' or cfg['autoText']):
                stage(jid,'AI改写剧名与简介',8);step('text',lambda:rewrite(p,cfg,jid))
            material=p
            continuing=kind=='pipeline' and bool(job.get('options',{}).get('continueFromJob'))
            if kind=='pipeline' and cfg['subtitleAuto'] and not continuing:
                stage(jid,'先去字幕，完成后进入重组分集',10)
                cleaned=repair_subtitles(p,cfg,jid,True)
                material={**p,'videos':video_entries(cleaned)}
                completed('subtitle')
            if kind in ['pipeline','transcode']:
                if continuing:
                    stage(jid,'复用已完成分集，继续后续制作',35)
                    files=continued_video_files(job['options'],pid)
                    checkpoint(jid,f'{pid}:video',files)
                else:
                    stage(jid,'使用去字幕结果重组分集' if material is not p else '整剧重组与画面处理',35)
                    files=transcode(material,cfg,jid)
                set_final(p,files,jid);completed('video')
            if kind=='subtitle':
                repair_subtitles(p,cfg,jid,job.get('options',{}).get('subtitleAllSelected',False))
                completed('subtitle')
            if kind=='preview':preview_subtitle(p,cfg,jid)
            if kind=='playback':playback_proxy(p,cfg,jid)
            if kind=='hook':promotion_hook(p,cfg,jid)
            if kind=='promotion_package':promotion_package(p,cfg,jid)
            if kind=='highlight' or (kind=='pipeline' and cfg['pipelineHighlight']):
                inputs=[] if kind=='highlight' and job.get('options',{}).get('highlightOriginalSelection') else final_videos(p)
                options=cfg
                if inputs:
                    material={**p,'videos':video_entries([a['path'] for a in inputs])}
                    options=cfg|{'speed':1,'cropPercent':0,'brightness':0,'contrast':1,'saturation':1,'markAI':False}
                else:material=p
                set_final(p,highlight(material,options,jid),jid)
                completed('highlight')
            if kind in ['cover','coverAI'] or (kind=='pipeline' and cfg['autoCover']):
                stage(jid,'制作封面',75);step('cover',lambda:cover(p,cfg,jid,ai=kind=='coverAI' or (kind=='pipeline' and cfg['autoAICover'])))
            if kind=='records' or (kind=='pipeline' and cfg['autoRecords']):
                stage(jid,'生成真实工程记录',82);step('records',lambda:records(p,cfg,jid))
            if kind=='cost' or (kind=='pipeline' and cfg['autoCost']):
                stage(jid,'生成成本报告',90);step('cost',lambda:costs(p,cfg,jid))
            if kind=='package':stage(jid,'整理发布材料',80);step('package',lambda:package(p,cfg,jid))
            if kind=='upload':upload_videos(p,cfg,jid)
            if kind in ['pipeline','transcode']:p['status']='已处理' if len(p['videos'])==len(get_project(pid)['videos']) else '部分集已处理'
            stage(jid,f'{p["name"]}：完成 ({index+1}/{len(ids)})',100)
            checkpoint(jid,f'{pid}:project-done',[a['path'] for a in p['artifacts'] if a['job']==jid and Path(a['path']).is_file()])
            save()
        job.update(status='completed',progress=100,stage='全部完成')
        log('任务完成',job=jid)
    except InterruptedError:
        job.update(status='paused' if job.get('pauseRequested') else 'cancelled',stage='已暂停，可从检查点继续' if job.get('pauseRequested') else '已停止；已完成文件保留')
        log('任务已停止',level='warning',job=jid)
    except Exception as e:
        msg=ai.redact(str(e))
        for secret in SECRETS.values():
            if secret:msg=msg.replace(secret,'[已隐藏]')
        job.update(status='failed',stage='任务失败',error=msg)
        log(msg,level='error',job=jid)
    finally:
        connections.TRACE_CALLBACKS.pop(jid,None)
        telemetry.settle(job,job['status'])
        job['runSeconds']=job.get('runSeconds',0)+time.time()-run_start
        job['finished']=time.time();save()

class SelectedProject(dict):
    """A fixed video selection while project metadata and artifacts still persist."""
    def __init__(self,base,paths):
        dict.__init__(self,base);self.base=base
        mapping={v['path']:v for v in base['videos']}
        dict.__setitem__(self,'videos',[copy.deepcopy(mapping[p]) for p in paths])
    def __setitem__(self,key,value):
        dict.__setitem__(self,key,value)
        if key!='videos':self.base[key]=value
    def __getitem__(self,key):
        return dict.__getitem__(self,key) if key=='videos' else self.base[key]
    def get(self,key,default=None):
        return dict.get(self,key,default) if key=='videos' else self.base.get(key,default)
    def __contains__(self,key):
        return key=='videos' or key in self.base
    def setdefault(self,key,default=None):
        if key not in self:self[key]=default
        return self[key]
    def update(self,*args,**kwargs):
        for key,value in dict(*args,**kwargs).items():self[key]=value

def source_snapshot(ids,selections=None):
    def videos(pid):
        entries={v['path']:v for v in get_project(pid)['videos']}
        if selections and pid in selections:
            if any(p not in entries for p in selections[pid]):raise ValueError('所选源视频已移除，请核对后重新创建任务')
            return [entries[p] for p in selections[pid]]
        return list(entries.values())
    return {pid:[{'path':v['path'],'bytes':Path(v['path']).stat().st_size,'mtime':Path(v['path']).stat().st_mtime_ns}
                 for v in videos(pid)] for pid in ids}

def ensure_runner():
    global QUEUE_THREAD
    with LOCK:
        if QUEUE_STOP.is_set():raise ValueError('服务正在关闭，请重新打开应用后继续')
        if QUEUE_THREAD is None or not QUEUE_THREAD.is_alive():
            QUEUE_THREAD=threading.Thread(target=queue_runner,daemon=True);QUEUE_THREAD.start()
        QUEUE_WAKE.set()

def queue_runner():
    while not QUEUE_STOP.is_set():
        with LOCK:
            pending=sorted((j for j in JOBS.values() if j['status']=='queued'),key=lambda j:j.get('queueAt',j['created']))
            job=pending[0] if pending else None
            if job:job['status']='running';persist_jobs()
        if job:worker(job['id'],job['ids'],job['kind'],job['config'])
        else:QUEUE_WAKE.wait(1);QUEUE_WAKE.clear()

def control_job(jid,action):
    with LOCK:
        job=JOBS.get(jid)
        if not job:raise ValueError('任务不存在')
        if action=='resume':
            if job['kind']=='video_generate':raise ValueError('视频生成功能已移除，历史任务不能恢复；已有视频仍可使用')
            if job['kind'] not in LOCAL_JOB_KINDS:raise ValueError('登录或上传任务不能自动恢复，请在发布页重新核对并发起')
            if job['status'] not in ['paused','interrupted','failed','cancelled']:raise ValueError('该任务不能恢复')
            if job.get('sourceSnapshot')!=source_snapshot(job['ids'],job.get('options',{}).get('sourceVideoPaths')):raise ValueError('源素材已改变，请创建新任务；未复用旧检查点')
            job['cancel'].clear();job.update(status='queued',stage='已加入恢复队列',pauseRequested=False,queueAt=time.time(),attempts=job.get('attempts',1)+1)
            job.pop('error',None);job.pop('finished',None);persist_jobs();ensure_runner()
        elif action in ['pause','cancel']:
            if job['status'] in ['completed']:raise ValueError('任务已完成')
            if action=='pause' and job['kind'] not in LOCAL_JOB_KINDS:raise ValueError('登录或上传任务请停止后重新核对')
            job['pauseRequested']=action=='pause';job['cancel'].set()
            if job['status'] not in ['running','pausing']:
                job.update(status='paused' if action=='pause' else 'cancelled',stage='已暂停' if action=='pause' else '已取消',finished=time.time())
            else:
                job.update(status='pausing' if action=='pause' else 'running',stage='正在暂停，保留已完成视频' if action=='pause' else '正在停止')
                for proc in list(job['processes']):
                    try:proc.kill()
                    except OSError:pass
            persist_jobs()
        else:raise ValueError('任务操作无效')
    return {'id':jid,'status':job['status']}

def shutdown_jobs():
    QUEUE_STOP.set();QUEUE_WAKE.set()
    with LOCK:
        for job in JOBS.values():
            if job['status'] in ['running','queued','pausing']:
                job['pauseRequested']=True;job['cancel'].set()
                job.update(status='interrupted',stage='应用已关闭，重新打开后可继续')
                for proc in list(job['processes']):
                    try:proc.kill()
                    except OSError:pass
        persist_jobs()
    vsr_pool.POOL.close()
    lama_local.POOL.close()
    if PUBLISH_WORKSPACE is not None:PUBLISH_WORKSPACE.shutdown()

def hook_inputs(p,kind):
    if kind=='original':return p['videos']
    if kind=='final':return final_videos(p)
    if kind=='subtitle':return [a for a in p['artifacts'] if a['kind'] in ('AI字幕处理','字幕处理') and Path(a['path']).is_file()]
    raise ValueError('请选择原片、去字幕结果或当前成片')

def job_preflight(body):
    """Validate selection and storage without queuing or contacting model providers."""
    cfg=validate_config(copy.deepcopy(STATE['config']));kind=body.get('kind','pipeline')
    if 'aiSelection' in body:cfg['aiSelection']=copy.deepcopy(body['aiSelection']);connections.validate(cfg)
    if kind not in ('pipeline','transcode','subtitle','highlight'):raise ValueError('此制作类型不支持选集预检')
    ids=body.get('ids',[]);requested=body.get('sourceVideoPaths')
    if not isinstance(ids,list) or not ids or not isinstance(requested,dict) or set(requested)!=set(ids):raise ValueError('请核对任务剧集和所选视频')
    videos=[]
    for pid in ids:
        p=get_project(pid);available={v['path']:v for v in p['videos']};paths=requested[pid]
        if not isinstance(paths,list) or not paths or len(set(paths))!=len(paths) or any(x not in available for x in paths):raise ValueError('请选择属于该剧集的有效本地视频')
        videos.extend(available[x] for x in paths)
    continuing=bool(body.get('continueFromJob'))
    errors=[];warnings=[];uses_subtitle=kind=='subtitle' or (kind=='pipeline' and cfg['subtitleAuto'] and not continuing)
    if continuing and not any(cfg[k] for k in ('autoText','pipelineHighlight','autoCover','autoRecords','autoCost')):errors.append('请至少开启一个后续制作步骤，再继续已完成分集')
    if kind=='pipeline' and cfg['autoCost'] and sum(float(row['amount']) for row in cfg['costRows'])<=0:errors.append('已开启成本报告，请先填写实际成本或关闭该步骤')
    for v in videos:
        if not Path(v['path']).is_file():errors.append('源视频已移动或删除：'+v['name'])
    if uses_subtitle and cfg['subtitleMode']=='vsr':
        if not (Path(cfg['vsrRoot'])/'backend/main.py').is_file() or not Path(cfg['vsrPython']).is_file():errors.append('本地AI去字幕引擎未配置或不完整，请到AI去字幕页核对')
        if cfg['vsrMode']=='lama-box' and not lama_local.ready(cfg['vsrRoot'],cfg['vsrPython']):errors.append('LaMa 本地模型未安装：需要 big-lama/big-lama.pt，可先切回 STTN')
    interfaces=[];channels=[]
    if uses_subtitle and cfg['subtitleMode']=='cloud':channels.append('subtitle')
    if kind=='pipeline' and cfg['autoText']:channels.append('text')
    if kind=='pipeline' and cfg['autoCover'] and cfg['autoAICover']:channels.append('image')
    for channel in channels:
        label={'text':'AI文案','image':'AI封面','subtitle':'云端去字幕'}[channel]
        try:
            choices=connections.candidates(cfg,channel)
            if choices:
                interfaces.append(label+'：'+' → '.join(c['name'] for c in choices))
                if not any(ai.get_key(connections.bound(cfg,c),channel) for c in choices):errors.append(label+'接口缺少本次会话密钥，请到AI接口管理填写')
            elif channel=='subtitle':
                interfaces.append(label+'：原配置接口')
                if not cfg['subtitleEndpoint']:errors.append('云端去字幕接口未配置')
            else:
                interfaces.append(label+'：原配置 '+ai.connection(cfg,channel)['model'])
                if not ai.get_key(cfg,channel):errors.append(label+'接口缺少本次会话密钥')
        except (ValueError,ai.ProviderError) as e:errors.append(label+'：'+str(e))
    multiplier=cfg['highlightCount']*(1+(cfg['highlightMax'] if cfg['highlightRandom'] else cfg['highlightFixed']))*2 if kind=='highlight' else (3 if uses_subtitle else 2)
    output=Path(cfg['output']).expanduser().resolve();estimate=max(256*1024**2,sum(v.get('bytes',0) for v in videos)*multiplier)
    try:
        for pid in ids:project_dir(get_project(pid),cfg)
        check_storage(output,estimate)
        free=shutil.disk_usage(output).free
    except (OSError,ValueError) as e:errors.append(str(e));free=None
    warnings.append('空间按源文件体积估算；编码、画幅及模型缓存会影响实际用量')
    return {'ok':not errors,'errors':errors,'warnings':warnings,'count':len(videos),'files':[v['name'] for v in videos],
            'output':str(output),'estimatedBytes':estimate,'freeBytes':free,'subtitle':uses_subtitle,
            'subtitleMode':cfg['subtitleMode'],'subtitleRegions':subtitle_regions(cfg) if uses_subtitle else [],
            'width':cfg['width'],'height':cfg['height'],'regroup':cfg['regroup'],'interfaces':interfaces,
            'workflowPlan':production_plan(kind,cfg,continuing)}

def create_job(body):
    licensing.require(DATA)
    kind=body.get('kind','pipeline')
    if kind not in LOCAL_JOB_KINDS|{'login','upload'}:
        raise ValueError('未知任务类型')
    ids=body.get('ids',[])
    if not ids and kind!='login':raise ValueError('请先勾选至少一部剧')
    account_ids=body.get('accountIds',[])
    if kind in ['upload','login']:
        if not account_ids:raise ValueError('请至少选择一个账号')
        if len(account_ids)!=len(set(account_ids)):raise ValueError('账号重复')
        if any(not any(a['id']==aid for a in STATE['accounts']) for aid in account_ids):raise ValueError('账号不存在')
        if kind=='login' and len(account_ids)!=1:raise ValueError('一次只能登录一个账号')
        if kind=='upload' and body.get('confirmUpload') is not True:raise ValueError('请在上传确认窗口核对材料、账号和模式')
    if len(ids)!=len(set(ids)):raise ValueError('任务中存在重复剧目')
    cfg=validate_config(copy.deepcopy(STATE['config']))
    if 'aiSelection' in body:
        cfg['aiSelection']=copy.deepcopy(body['aiSelection']);connections.validate(cfg)
    for pid in ids:get_project(pid)
    options={'encodingPlanVersion':2} if kind in ('pipeline','transcode') else {}
    parent=None
    if 'continueFromJob' in body:
        if not isinstance(body['continueFromJob'],str) or not body['continueFromJob']:raise ValueError('接续任务编号无效')
        parent=JOBS.get(body['continueFromJob'])
        if kind!='pipeline' or not parent or parent['kind']!='transcode' or parent['status']!='completed' or ids!=parent['ids']:
            raise ValueError('只能从同一组剧集已完成的合并切分任务继续')
        paths=parent.get('options',{}).get('sourceVideoPaths') or {pid:[v['path'] for v in get_project(pid)['videos']] for pid in ids}
        if parent.get('sourceSnapshot')!=source_snapshot(ids,paths):raise ValueError('源素材发生变化，请核对后重新制作')
        if 'sourceVideoPaths' in body and body['sourceVideoPaths']!=paths:raise ValueError('接续制作必须使用原任务的选集范围和顺序')
        body=body|{'sourceVideoPaths':copy.deepcopy(paths)}
        options['continueFromJob']=parent['id'];options['continuedVideoFiles']={}
        for pid in ids:
            files=cached(parent['id'],f'{pid}:video')
            if not files:raise ValueError('原分集文件或检查点已失效，不能继续后续制作')
            options['continuedVideoFiles'][pid]=[{'path':str(path),'bytes':path.stat().st_size,'mtime':path.stat().st_mtime_ns} for path in files]
            continued_video_files(options,pid)
    if kind=='jianying_project':
        options,_=jianying_request(body)
        if not jianying_projects.status(options['jianying'])['available']:raise ValueError('剪映草稿组件未安装')
        with LOCK:STATE['jianyingOptions']=copy.deepcopy(options['jianying']);save()
    elif kind=='track_sheet':
        options,_=track_request(body)
        plan=track_plan(body)
        if not plan['ok']:raise ValueError('；'.join(plan['errors']))
        with LOCK:STATE['trackSheetOptions']=copy.deepcopy(options['trackSheet']);save()
    elif kind in ('pipeline','transcode','subtitle','highlight'):
        requested=body.get('sourceVideoPaths')
        if requested is not None and (not isinstance(requested,dict) or set(requested)!=set(ids)):raise ValueError('所选源视频必须与任务剧集对应')
        fixed={}
        for pid in ids:
            available={v['path']:v for v in get_project(pid)['videos']}
            paths=requested[pid] if requested is not None else list(available)
            if not isinstance(paths,list) or not paths or any(not isinstance(x,str) or x not in available for x in paths):raise ValueError('请选择属于该剧集的有效本地视频')
            if len(paths)!=len(set(paths)):raise ValueError('制作集数重复')
            fixed[pid]=list(paths)
        options['sourceVideoPaths']=fixed
        if kind=='highlight':options['highlightOriginalSelection']=requested is not None
        # New standalone tasks cover every selected file in full. Persist this
        # option in the job so historical first-N tasks retain their own scope.
        if kind=='subtitle':options['subtitleAllSelected']=True
        if requested is not None or kind in ('pipeline','subtitle'):
            preflight=job_preflight(body|{'sourceVideoPaths':fixed})
            if not preflight['ok']:raise ValueError('；'.join(preflight['errors']))
    elif 'sourceVideoPaths' in body:raise ValueError('此制作类型不支持指定源视频集数')
    if kind=='highlight':highlight_render.validate(cfg,True)
    if kind=='hook':
        if len(ids)!=1:raise ValueError('一次请选择一部剧制作钩子')
        p=get_project(ids[0])
        if not p['rights']:raise ValueError('请先确认该剧拥有制作与发布授权')
        index=body.get('videoIndex',0)
        if type(index) is not int or index<0:raise ValueError('请选择有效的源视频')
        source_kind=body.get('sourceKind','original');inputs=hook_inputs(p,source_kind)
        if source_kind=='original' and 'sourcePath' not in body:
            if index>=len(p['videos']):raise ValueError('请选择有效的源视频')
            v=p['videos'][index]
        else:
            entry=next((v for v in inputs if v['path']==body.get('sourcePath')),None)
            if not entry:raise ValueError('所选版本的视频不存在或不属于该剧，请刷新后选择')
            v={**entry,**probe(entry['path'])}
        start=number(body.get('start',0),0,max(0,v['duration']-1),'钩子起点')
        seconds=number(body.get('seconds',min(30,v['duration']-start)),1,min(180,v['duration']-start),'钩子时长')
        options={'videoIndex':index,'start':start,'seconds':seconds,'title':promotions.text(body.get('title',p['name']+' · 高光钩子'),'钩子标题',100,True),
            'cta':promotions.text(body.get('cta',''),'结尾引导文字',80)}
        st=Path(v['path']).stat();options.update(hookSource=copy.deepcopy(v),hookSourceKind=source_kind,hookSourceStat={'bytes':st.st_size,'mtime':st.st_mtime_ns})
    if kind=='promotion_package':
        if len(ids)!=1:raise ValueError('一次请选择一部剧导出推广任务')
        row=promotions.task(STATE,body.get('promotionTaskId'))
        if row['projectId']!=ids[0] or row['status']=='cancelled':raise ValueError('推广任务与剧集不匹配或已取消')
        if not get_project(ids[0])['rights']:raise ValueError('请先确认素材授权')
        options={'promotionTask':copy.deepcopy(row)}
    if kind in ['preview','playback']:
        if len(ids)!=1:raise ValueError('一次只试处理一部剧的视频')
        index=body.get('videoIndex',0)
        if not isinstance(index,int) or not 0<=index<len(get_project(ids[0])['videos']):raise ValueError('请选择有效源视频')
        options={'videoIndex':index}
        if kind=='preview':
            video=get_project(ids[0])['videos'][index];start=number(body.get('start',0),0,max(0,video['duration']-.25),'试处理起点')
            seconds=body.get('seconds',5)
            if seconds not in [5,10]:raise ValueError('试处理时长请选择5秒或10秒')
            options.update(start=start,seconds=min(seconds,video['duration']-start))
    with LOCK:
        if kind in ['login','upload'] and any(j['status'] in ['queued','running','pausing'] for j in JOBS.values()):raise ValueError('请先完成制作队列，再核对登录或上传任务')
        if parent:
            existing=next((j for j in JOBS.values() if j.get('options',{}).get('continueFromJob')==parent['id'] and j['status']!='cancelled'),None)
            if existing:return {'id':existing['id'],'reused':True}
        jid=uuid.uuid4().hex[:12]
        JOBS[jid]={'id':jid,'kind':kind,'ids':ids,'status':'queued','stage':'准备中','progress':0,'created':time.time(),
                   'accountIds':account_ids,'completedStages':[],'cancel':threading.Event(),'processes':set(),
                   'config':cfg,'options':options,'sourceSnapshot':source_snapshot(ids,options.get('sourceVideoPaths')),'checkpoints':{},'attempts':1,'queueAt':time.time()}
        persist_jobs()
    ensure_runner()
    return {'id':jid}

def import_source_video(directory,meta,video,cover,details):
    """Register only a verified completed download; preserve the original source facts."""
    pid=hashlib.sha256(str(directory).casefold().encode()).hexdigest()[:16]
    project={'id':pid,'source':str(directory),'name':meta['title'],'originalName':meta['title'],
        'intro':meta['intro'],'originalIntro':meta['intro'],'cover':str(cover) if cover else '',
        'videos':[{'path':str(video),'name':video.name,'bytes':video.stat().st_size,'mtime':video.stat().st_mtime,**details}],
        'status':'未处理','artifacts':[],'review':'待准备','rights':False,'notes':'按所选集数导入'+('原生来源' if meta.get('platform')=='guoapp_hongguo' else '网页公开版')+'，非整剧完成声明',
        'sourceMetadata':copy.deepcopy(meta)}
    with LOCK:
        old=next((p for p in STATE['projects'] if p['id']==pid),None)
        if old:
            videos=[v for v in old['videos'] if v['path']!=str(video)]+project['videos']
            videos.sort(key=lambda v:(v.get('sourceEpisode',0),natural(v['name'])))
            old.update(videos=videos,sourceMetadata=copy.deepcopy(meta))
            if not old.get('cover') and cover:old['cover']=str(cover)
        else:STATE['projects'].append(project)
        save()
    return pid

def source_manager():
    global SOURCE_MANAGER
    with LOCK:
        if SOURCE_MANAGER is None:SOURCE_MANAGER=guoapp_source.Manager(DATA,import_source_video,probe)
        return SOURCE_MANAGER

def public_state():
    with LOCK:
        result=copy.deepcopy(STATE)
        result['jobs']=[{**{k:v for k,v in j.items() if k not in ['cancel','processes','config','sourceSnapshot','checkpoints']},
                         'workflowPlan':production_plan(j['kind'],j['config'],bool(j.get('options',{}).get('continueFromJob'))),
                         'telemetry':telemetry.snapshot(j),'completedProjects':sum(k.endswith(':project-done') for k in j.get('checkpoints',{})),
                         'checkpointCount':len(j.get('checkpoints',{})), 'resumable':j['kind'] in LOCAL_JOB_KINDS} for j in JOBS.values()][-200:]
        result['secrets']={k:bool(v) for k,v in SECRETS.items()}
        result['aiKeyStatus']=ai.key_status(result['config'])
        result['aiProviders']=ai.PRESETS
        result['connectionKeyStatus']=connections.key_status(result['config'])
        result['promotionSummary']=promotions.summary(STATE)
        result['promotionEntryTypes']=promotions.TYPES
        return result


def publish_workspace():
    global PUBLISH_WORKSPACE
    with LOCK:
        if PUBLISH_WORKSPACE is None:
            def snapshot(pid):
                with LOCK:
                    project=copy.deepcopy(get_project(pid));config=copy.deepcopy(STATE['config'])
                    # Only the latest explicitly registered finished cut is eligible.
                    # Source videos and 5/10-second previews are never substituted.
                    if project.get('finalVideoPaths'):
                        entries={a['path']:a for a in project['artifacts']}
                        project['publishVideos']=[copy.deepcopy(entries.get(path,{'path':path})) for path in project['finalVideoPaths']]
                    else:
                        entries=[a for a in project['artifacts'] if a['kind']=='视频']
                        # Select the latest generation before checking existence.
                        # A missing latest output must not fall back to an older cut.
                        project['publishVideos']=copy.deepcopy([a for a in entries if a.get('job')==entries[-1].get('job')]) if entries else []
                    project['path']=project.get('publishFolder') or project.get('source','')
                    return project,config
            PUBLISH_WORKSPACE=publish_workflow.Workspace(STATE,DATA,LOCK,save,snapshot,probe)
        return PUBLISH_WORKSPACE

def allowed_file(path):
    path=Path(path).resolve()
    # 不允许借由素材文件服务读取脚本、配置或密钥；仅暴露明确注册的素材和生成文件。
    allowed=set()
    for p in STATE['projects']:
        allowed.update(v['path'] for v in p['videos'])
        if p['cover']:allowed.add(p['cover'])
        allowed.update(a['path'] for a in p['artifacts'])
    if str(path) not in allowed or not path.is_file():raise ValueError('文件不在素材清单或已被移动')
    return path

class Handler(BaseHTTPRequestHandler):
    protocol_version='HTTP/1.1'
    server_version='ManjuStudio'
    sys_version=''
    def end_headers(self):
        local_session.headers(self)
        super().end_headers()
    def handle(self):
        try:super().handle()
        except (ConnectionResetError,ConnectionAbortedError,BrokenPipeError):pass
    def log_message(self,*args):pass
    def reply(self,status=200,value=None):
        data=json.dumps(value if value is not None else {},ensure_ascii=False).encode()
        self.send_response(status);self.send_header('Content-Type','application/json; charset=utf-8')
        self.send_header('Content-Length',str(len(data)));self.send_header('Cache-Control','no-store')
        self.end_headers();self.wfile.write(data)
    def local_check(self):
        local_session.check(self)
    def session_check(self):
        session=getattr(self.server,'session',None)
        if (getattr(sys,'frozen',False) or session is not None) and (not session or not session.authenticated(self.headers)):raise PermissionError('本机应用会话无效，请从桌面快捷方式启动')
    def file(self,path,download=False):
        size=path.stat().st_size;start=0;end=size-1
        match=re.fullmatch(r'bytes=(\d*)-(\d*)',self.headers.get('Range',''))
        partial=False
        if match:
            a,b=match.groups();partial=True
            if a:start=int(a);end=min(int(b),end) if b else end
            elif b:start=max(0,size-int(b))
            else:partial=False
            if start> end or start>=size:
                self.send_response(416);self.send_header('Content-Range',f'bytes */{size}');self.send_header('Content-Length','0');self.end_headers();return
        self.send_response(206 if partial else 200)
        self.send_header('Content-Type',mimetypes.guess_type(path.name)[0] or 'application/octet-stream')
        self.send_header('Content-Length',str(end-start+1));self.send_header('Accept-Ranges','bytes')
        self.send_header('X-Content-Type-Options','nosniff');self.send_header('Cache-Control','no-cache')
        if partial:self.send_header('Content-Range',f'bytes {start}-{end}/{size}')
        if download:self.send_header('Content-Disposition',"attachment; filename*=UTF-8''"+urllib.parse.quote(path.name))
        self.end_headers()
        with path.open('rb') as f:
            f.seek(start);remaining=end-start+1
            while remaining>0:
                chunk=f.read(min(1024*1024,remaining))
                if not chunk:break
                self.wfile.write(chunk);remaining-=len(chunk)
    def do_GET(self):
        try:
            self.local_check();url=urllib.parse.urlsplit(self.path);q=urllib.parse.parse_qs(url.query)
            session=getattr(self.server,'session',None)
            if session and secrets.compare_digest(url.path,'/launch/'+session.launch):
                self.send_response(302);self.send_header('Location','/');self.send_header('Cache-Control','no-store')
                self.send_header('Set-Cookie','manju_session='+session.cookie+'; HttpOnly; SameSite=Strict; Path=/')
                self.send_header('Content-Length','0');self.end_headers();return
            self.session_check()
            if url.path=='/api/publish/workspace':
                licensing.require(DATA)
                return self.reply(value=publish_workspace().public())
            if url.path=='/api/publish/file':
                licensing.require(DATA)
                return self.file(publish_workspace().file(q.get('projectId',[''])[0],q.get('path',[''])[0]),q.get('download',['0'])[0]=='1')
            if url.path=='/api/license':return self.reply(value=licensing.status(DATA))
            if url.path=='/api/component':return self.reply(value={'progress':engine_manager.progress(),'active':engine_manager.active(DATA)})
            if not licensing.status(DATA)['valid']:
                if url.path in ['/','/index.html']:return self.file(PUBLIC/'activation.html')
                if url.path not in ['/activation.js','/activation.css']:raise PermissionError('请先激活本机授权')
            if url.path=='/api/state':return self.reply(value=public_state())
            if url.path=='/api/source/tasks':return self.reply(value={'tasks':source_manager().list(),'access':hongguo_source.ACCESS.snapshot(),'native':source_manager().core.status()})
            if url.path=='/api/source/report':
                value=source_manager().report(q.get('id',[''])[0]);data=json.dumps(value,ensure_ascii=False,indent=2).encode('utf-8')
                self.send_response(200);self.send_header('Content-Type','application/json; charset=utf-8');self.send_header('Content-Length',str(len(data)))
                self.send_header('Content-Disposition','attachment; filename="source-diagnostic.json"');self.end_headers();self.wfile.write(data);return
            if url.path=='/api/source/cover':
                data,mime=source_manager().cover(q.get('id',[''])[0],q.get('platform',[''])[0])
                self.send_response(200);self.send_header('Content-Type',mime);self.send_header('Content-Length',str(len(data)));self.end_headers();self.wfile.write(data);return
            if url.path=='/api/jobs/report':
                jid=urllib.parse.parse_qs(url.query).get('id',[''])[0]
                with LOCK:
                    if jid not in JOBS:raise ValueError('任务不存在')
                    return self.reply(value=telemetry.report(JOBS[jid]))
            if url.path=='/api/environment':return self.reply(value=environment(force=q.get('refresh',['0'])[0]=='1'))
            if url.path=='/api/environment/report':
                value=environment()
                return self.reply(value={'format':1,'product':'ManjuStudio','generatedAt':time.time(),
                    'ffmpeg':value['ffmpeg'],'gpu':value['gpu'],'cpu':value['cpu'],'cores':value['cores'],
                    'python':value['python'],'font':value['font'],
                    'note':'此诊断不含授权码、接口密钥、原始错误或本地文件路径；硬件编码实测不代表所有显卡型号均已验证。'})
            if url.path in ['/api/promotion/revenue-template','/api/promotion/revenue-export']:
                data=promotions.csv_export(STATE,url.path.endswith('template'))
                self.send_response(200);self.send_header('Content-Type','text/csv; charset=utf-8');self.send_header('Content-Length',str(len(data)))
                self.send_header('Content-Disposition','attachment; filename="promotion-revenue.csv"');self.end_headers();self.wfile.write(data);return
            if url.path=='/api/file':return self.file(allowed_file(q.get('path',[''])[0]),q.get('download',['0'])[0]=='1')
            if url.path=='/api/logs/download':
                data='\n'.join(f'{e["time"]} [{e["level"]}] {e["message"]}' for e in STATE['logs']).encode('utf-8-sig')
                self.send_response(200);self.send_header('Content-Type','text/plain; charset=utf-8');self.send_header('Content-Length',str(len(data)))
                self.send_header('Content-Disposition','attachment; filename="manju-logs.txt"');self.end_headers();self.wfile.write(data);return
            relative=urllib.parse.unquote(url.path).lstrip('/') or 'index.html';path=(PUBLIC/relative).resolve()
            if not path.is_relative_to(PUBLIC) or not path.is_file():return self.reply(404,{'error':'页面不存在'})
            self.file(path)
        except (BrokenPipeError,ConnectionResetError,ConnectionAbortedError):pass
        except PermissionError as e:self.reply(403,{'error':str(e)})
        except Exception as e:self.reply(400,{'error':str(e)})
    def do_POST(self):
        try:
            self.local_check()
            self.session_check()
            if self.headers.get('X-Local-Client')!='1':return self.reply(403,{'error':'缺少本地客户端请求头'})
            length=int(self.headers.get('Content-Length','0'))
            if length<0 or length>1024*1024 or self.headers.get('Transfer-Encoding'):raise ValueError('请求大小无效')
            if self.headers.get('Content-Type','').split(';')[0]!='application/json':raise ValueError('请求格式必须为 JSON')
            body=json.loads(self.rfile.read(length) or b'{}');path=urllib.parse.urlsplit(self.path).path
            if not isinstance(body,dict):raise ValueError('请求必须为对象')
            if path=='/api/license/activate':return self.reply(value=licensing.activate(DATA,body.get('token','')))
            licensing.require(DATA)
            if path=='/api/publish/inspect':return self.reply(value=publish_workspace().inspect(body.get('projectId',''),body.get('overrides')))
            if path=='/api/publish/settings':return self.reply(value=publish_workspace().settings(body.get('projectId',''),body.get('overrides',{})))
            if path=='/api/publish/discover':return self.reply(value=publish_workspace().discover(body.get('folder','')))
            if path=='/api/publish/import':return self.reply(value=publish_workspace().import_folders(body.get('folders',[])))
            if path=='/api/publish/export':return self.reply(value=publish_workspace().export(body.get('projectId','')))
            if path=='/api/publish/browser':return self.reply(value=publish_workspace().browser(body.get('accountId',''),body.get('action',''),body.get('projectId','')))
            if path=='/api/publish/record':return self.reply(value=publish_workspace().record(body))
            if path=='/api/publish/open':return self.reply(value=publish_workspace().open(body.get('projectId',''),body.get('taskId',''),body.get('key','')))
            if path=='/api/tracks/import':return self.reply(value=import_track_files(body.get('paths')))
            if path=='/api/jianying/status':return self.reply(value=jianying_projects.status(body.get('jianying',STATE.get('jianyingOptions',{}))))
            if path=='/api/jianying/plan':
                options,videos=jianying_request(body)
                return self.reply(value={'count':len(videos),'duration':sum(v['duration'] for v in videos),'sourceKind':options['sourceKind'],'editorValidated':False})
            if path=='/api/jianying/open':
                job=JOBS.get(body.get('id'));index=body.get('index',0)
                if not job or job['kind']!='jianying_project' or type(index) is not int or not 0<=index<len(job.get('jianyingResults',[])):raise ValueError('没有对应的剪映工程')
                target=Path(job['jianyingResults'][index]['directory']);jianying_projects.check(target)
                if os.name=='nt':os.startfile(str(target))
                return self.reply(value={'directory':str(target)})
            if path=='/api/jianying/relocate':return self.reply(value=jianying_projects.relocate(body.get('directory','')))
            if path=='/api/highlight/templates':
                cfg=validate_config(STATE['config']|{'highlightTailDirectory':body.get('directory',STATE['config']['highlightTailDirectory'])})
                return self.reply(value={'files':[p.name for p in highlight_render.templates(cfg)]})
            if path=='/api/tracks/plan':return self.reply(value=track_plan(body))
            if path=='/api/tracks/options':
                with LOCK:STATE['trackSheetOptions']=track_sheets.options(body);save()
                return self.reply(value={'options':STATE['trackSheetOptions']})
            if path=='/api/tracks/devices':
                devices=[]
                try:
                    proc=subprocess.run(['nvidia-smi','--query-gpu=index,name','--format=csv,noheader'],capture_output=True,timeout=5,creationflags=getattr(subprocess,'CREATE_NO_WINDOW',0))
                    for line in proc.stdout.decode('utf-8','replace').splitlines():
                        index,sep,name=line.partition(',')
                        if sep and index.strip().isdigit() and int(index.strip())<32:devices.append({'index':int(index.strip()),'name':name.strip()})
                except (OSError,subprocess.TimeoutExpired):pass
                return self.reply(value={'devices':devices,'note':'任务按素材实测硬件解码；镜头分析和PNG绘图使用CPU'})
            if path=='/api/tracks/open':
                job=JOBS.get(body.get('id'));index=body.get('index',0)
                if not job or job['kind']!='track_sheet' or type(index) is not int or not 0<=index<len(job.get('trackResults',[])):raise ValueError('没有对应的工程图结果')
                target=allowed_file(job['trackResults'][index]['path']).parent
                if os.name=='nt':os.startfile(str(target))
                return self.reply(value={'path':str(target)})
            if path=='/api/tracks/export':
                job=JOBS.get(body.get('id'))
                if not job or job['kind']!='track_sheet' or not job.get('trackResults'):raise ValueError('请先生成工程图')
                if job['status'] in ['running','queued','pausing']:raise ValueError('请等待当前任务结束后导出')
                rows=copy.deepcopy(job['trackResults']);directory=DATA/'exports';directory.mkdir(exist_ok=True)
                target=directory/('分轨工程图_'+job['id']+'_'+uuid.uuid4().hex[:8]+'.zip');temp=target.with_suffix('.part')
                try:
                    with zipfile.ZipFile(temp,'w',compression=zipfile.ZIP_STORED) as archive:
                        for i,row in enumerate(rows):
                            image=allowed_file(row['path']);record=allowed_file(row['manifest'])
                            if sha256(image)!=row['imageSHA256'] or sha256(record)!=row.get('manifestSHA256'):raise ValueError('工程图或分析记录已修改，请重新生成后导出')
                            archive.write(image,f'{i+1:02d}_{image.name}');archive.write(record,f'{i+1:02d}_{record.name}')
                    temp.replace(target);register(get_project(rows[0]['projectId']),target,'工程图压缩包',job['id'])
                finally:temp.unlink(missing_ok=True)
                return self.reply(value={'path':str(target),'count':len(rows)})
            if path=='/api/source/resolve':
                value=source_manager().resolve(body.get('text',''),body.get('refresh') is True);return self.reply(value=value)
            if path=='/api/source/start':return self.reply(value=source_manager().start(body))
            if path=='/api/jobs/preflight':return self.reply(value=job_preflight(body))
            if path=='/api/promotion/inputs':
                p=get_project(body.get('projectId'));entries=hook_inputs(p,body.get('sourceKind','original'))
                return self.reply(value={'videos':[{**e,**probe(e['path'])} for e in entries]})
            if path=='/api/source/cancel':return self.reply(value=source_manager().cancel(body.get('id','')))
            if path=='/api/source/control':return self.reply(value=source_manager().control(body.get('id',''),body.get('action',''),body.get('episodes')))
            if path=='/api/source/access-resume':
                hongguo_source.ACCESS.resume()
                with source_manager().lock:source_manager().start_worker()
                return self.reply(value={'access':hongguo_source.ACCESS.snapshot()})
            if path=='/api/source/open':
                row=next((r for r in source_manager().list() if r['id']==body.get('id')),None)
                target=Path(row['savedDirectory']) if row and row.get('savedDirectory') else None
                if not target or not target.is_dir():raise ValueError('该任务还没有保存目录')
                if os.name=='nt':os.startfile(str(target))
                return self.reply(value={'path':str(target)})
            if path.startswith('/api/promotion/'):
                with LOCK:
                    if path=='/api/promotion/entry':
                        p=get_project(body.get('projectId'));p['promotionEntry']=promotions.entry(body.get('entry'))
                        p.setdefault('promotionEntries',{})[p['promotionEntry']['type']]=copy.deepcopy(p['promotionEntry']);save();return self.reply(value={'entry':p['promotionEntry']})
                    if path=='/api/promotion/task':
                        result=promotions.make_task(STATE,get_project(body.get('projectId')),body);save();return self.reply(value=result)
                    if path=='/api/promotion/receipt':
                        result=promotions.record(STATE,body);save();return self.reply(value=result)
                    if path=='/api/promotion/cancel':
                        row=promotions.task(STATE,body.get('taskId'))
                        if row['status']=='submitted':raise ValueError('已人工确认发布的任务保留记录，不能直接取消')
                        row.update(status='cancelled',mountStatus='pending',updatedAt=time.time());save();return self.reply(value=row)
                    if path=='/api/promotion/dispatch':
                        row=promotions.task(STATE,body.get('taskId'))
                        if row['status'] in ['submitted','cancelled']:raise ValueError('该任务已提交或取消，请核对记录，避免重复发布')
                        row.update(status='awaiting_manual',updatedAt=time.time());save()
                        return self.reply(value={'url':row['entry']['backendUrl'],'task':row,'message':'请在平台选择挂载入口并发布，然后回填核对结果'})
                    if path=='/api/promotion/revenue':
                        result=promotions.save_revenues(STATE,[body]);save();return self.reply(value=result)
                    if path=='/api/promotion/revenue-import':
                        result=promotions.import_csv(STATE,body.get('csv',''));save();return self.reply(value=result)
                return self.reply(404,{'error':'推广接口不存在'})
            if path=='/api/component/import':
                def apply_engine(value):
                    with LOCK:
                        STATE['config'].update(vsrRoot=value['root'],vsrPython=value['python'],vsrMode=value['verifiedMode'],subtitleMode='vsr');save()
                engine_manager.start(body.get('path',''),DATA,apply_engine)
                return self.reply(value={'started':True})
            if path=='/api/config':
                with LOCK:
                    cfg=validate_config(STATE['config']|{k:v for k,v in body.items() if k in DEFAULT})
                    STATE['config']=cfg;save()
                return self.reply(value={'config':cfg})
            if path=='/api/secrets':
                for key in SECRETS:
                    if key in body:
                        if key=='apiKey':ai.set_key(STATE['config'],'text',str(body[key]))
                        else:SECRETS[key]=str(body[key])
                return self.reply(value={'saved':True,'message':'密钥仅在当前服务进程中保存，重启后需重新输入'})
            if path=='/api/ai/key':
                channel=body.get('channel','text')
                if body.get('copyText'):ai.copy_text_key_to_image(STATE['config'])
                else:ai.set_key(STATE['config'],channel,body.get('value',''))
                return self.reply(value={'saved':True,'status':ai.key_status(STATE['config'])})
            if path=='/api/ai/models':
                return self.reply(value={'models':ai.list_models(STATE['config'],body.get('channel','text'))})
            if path=='/api/ai/connections':
                with LOCK:
                    cfg=copy.deepcopy(STATE['config']);items=cfg['aiConnections'];action=body.get('action','upsert');cid=body.get('id') or body.get('connection',{}).get('id')
                    old=next((c for c in items if c['id']==cid),None)
                    if action=='delete':
                        if not old:raise ValueError('连接不存在')
                        if any(j['status'] in ['queued','running','pausing'] and any(c['id']==cid for c in j['config'].get('aiConnections',[])) for j in JOBS.values()):raise ValueError('队列中有使用这份接口配置的任务，请先暂停或停止后再移除')
                        cfg['aiConnections']=[c for c in items if c['id']!=cid]
                        for channel,rule in cfg['aiRouting'].items():
                            if rule['primary']==cid:rule['primary']=''
                            rule['backups']=[x for x in rule['backups'] if x!=cid]
                            if cfg['aiSelection'][channel]==cid:cfg['aiSelection'][channel]='auto'
                        ai.set_key(connections.bound(cfg,old),old['channel'],'')
                    elif action=='upsert':
                        entry=copy.deepcopy(body.get('connection',{}));entry['id']=cid or uuid.uuid4().hex[:16]
                        cfg['aiConnections']=[c for c in items if c['id']!=entry['id']]+[entry]
                        connections.validate(cfg)
                        saved=next(c for c in cfg['aiConnections'] if c['id']==entry['id'])
                        if old and old['channel']!=saved['channel']:raise ValueError('编辑时不能改变连接用途，请另建连接')
                        rule=cfg['aiRouting'][saved['channel']]
                        if not rule['primary']:rule['primary']=saved['id']
                        cid=saved['id']
                    else:raise ValueError('连接操作无效')
                    STATE['config']=validate_config(cfg);STATE.setdefault('aiConnectionTests',{}).pop(cid,None);save()
                return self.reply(value={'id':cid,'config':STATE['config']})
            if path in ['/api/ai/connection-key','/api/ai/connection-test','/api/ai/connection-models']:
                cfg=copy.deepcopy(STATE['config']);c=next((c for c in cfg['aiConnections'] if c['id']==body.get('id')),None)
                if c is None:raise ValueError('连接不存在')
                selected=connections.bound(cfg,c)
                if path.endswith('connection-key'):
                    ai.set_key(selected,c['channel'],body.get('value',''))
                    with LOCK:STATE.setdefault('aiConnectionTests',{}).pop(c['id'],None);save()
                    return self.reply(value={'saved':True,'hasKey':bool(ai.get_key(selected,c['channel']))})
                if path.endswith('connection-models'):
                    return self.reply(value={'models':ai.list_models(selected,c['channel'])})
                started=time.monotonic()
                try:
                    connections.require_key(cfg,c)
                    if c['channel']=='text':
                        result=ai.text_generate(selected,'你是连接测试助手。','请仅回复OK。',64)
                        result={'ok':True,'method':'text_call','output':ai.redact(result['text']),'model':ai.redact(result['model'])}
                    elif c['channel']=='image':
                        models=ai.list_models(selected,c['channel']);result={'ok':True,'method':'model_list','models':models[:20],'modelFound':c['model'] in models,'note':'仅验证模型列表，不代表已实际生成图片'}
                    else:raise ValueError('字幕适配器请用5秒样片试处理，接口没有统一无素材测试契约')
                    result.update(seconds=round(time.monotonic()-started,3),testedAt=time.time())
                except Exception as error:
                    with LOCK:STATE.setdefault('aiConnectionTests',{})[c['id']]={'ok':False,'error':ai.redact(str(error)),'testedAt':time.time()};save()
                    raise
                with LOCK:STATE.setdefault('aiConnectionTests',{})[c['id']]=result;save()
                return self.reply(value=result)
            if path=='/api/scan':return self.reply(value=scan(body.get('source','')))
            if path=='/api/project':
                with LOCK:
                    p=get_project(body['id'])
                    if any(p['id'] in j['ids'] and j['status'] in ['queued','running','pausing'] for j in JOBS.values()):
                        raise ValueError('该剧正在处理，完成或停止任务后再修改剧集信息')
                    if 'name' in body and not str(body['name']).strip():raise ValueError('剧名不能为空')
                    for key in ['name','intro','notes','review','rights']:
                        if key in body:
                            if key=='rights':p[key]=bool(body[key])
                            elif isinstance(body[key],str):p[key]=body[key][:20000]
                    if not p['name'].strip():raise ValueError('剧名不能为空')
                    if p['review'] not in ['待准备','待提交','已提交','审核通过','需修改']:raise ValueError('审核状态无效')
                    save()
                return self.reply(value=p)
            if path=='/api/remove':
                with LOCK:
                    pid=body['id']
                    if any(pid in j['ids'] and j['status'] in ['running','queued','pausing'] for j in JOBS.values()):raise ValueError('任务运行中，不能移除')
                    STATE['projects']=[p for p in STATE['projects'] if p['id']!=pid];save()
                return self.reply(value={'removed':True})
            if path=='/api/jobs':return self.reply(value=create_job(body))
            if path=='/api/jobs/control':return self.reply(value=control_job(body['id'],body.get('action','')))
            if path=='/api/cancel':
                result=control_job(body['id'],'cancel')
                return self.reply(value={**result,'cancelled':True})
            if path=='/api/test-api':
                cfg=copy.deepcopy(STATE['config'])
                response=ai.text_generate(cfg,'你是连接测试助手，请遵照用户要求简短回答。','请仅回复OK。',max_tokens=64)
                result={k:v for k,v in response.items() if k!='text'}
                result.update(ok=True,output=ai.redact(response['text']),message='文字模型实际调用成功；生图功能需要独立模型和密钥。')
                log(f'{response["provider"]} 文字模型测试成功，耗时 {response["seconds"]} 秒')
                return self.reply(value=result)
            if path=='/api/accounts':
                name=str(body.get('name','')).strip()
                if not name:raise ValueError('请填写账号备注名')
                item={'id':uuid.uuid4().hex[:10],'name':name,'channelId':str(body.get('channelId','')),'type':body.get('type','剧集提审')}
                with LOCK:STATE['accounts'].append(item);save()
                return self.reply(value=item)
            if path=='/api/account-open':
                account=next((a for a in STATE['accounts'] if a['id']==body['id']),None)
                if not account:raise ValueError('账号备注不存在')
                candidates=[Path(os.environ.get('PROGRAMFILES(X86)','C:/Program Files (x86)'))/'Microsoft/Edge/Application/msedge.exe',
                            Path(os.environ.get('PROGRAMFILES','C:/Program Files'))/'Google/Chrome/Application/chrome.exe']
                browser=next((p for p in candidates if p.is_file()),None)
                if not browser:raise ValueError('未找到Edge或Chrome浏览器，无法创建独立账号窗口')
                profile=DATA/'accounts'/account['id'];profile.mkdir(parents=True,exist_ok=True)
                url='https://channels.weixin.qq.com/mcn/login' if account['type']=='剧集提审' else 'https://channels.weixin.qq.com/platform'
                subprocess.Popen([str(browser),'--user-data-dir='+str(profile),'--app='+url,'--no-first-run'],
                    creationflags=getattr(subprocess,'CREATE_NO_WINDOW',0))
                return self.reply(value={'opened':True})
            if path=='/api/account-remove':
                with LOCK:
                    if any(body['id'] in j['accountIds'] and j['status'] in ['queued','running'] for j in JOBS.values()):raise ValueError('该账号有运行中的任务')
                    STATE['accounts']=[a for a in STATE['accounts'] if a['id']!=body['id']];save()
                return self.reply(value={'removed':True})
            if path=='/api/browse':
                mode=body.get('mode','folder')
                if mode in ('images','documents'):
                    file_filter='Images|*.jpg;*.jpeg;*.png;*.webp' if mode=='images' else 'Reports|*.pdf;*.docx;*.csv;*.jpg;*.jpeg;*.png;*.webp'
                    script="Add-Type -AssemblyName System.Windows.Forms; $d=New-Object System.Windows.Forms.OpenFileDialog; $d.Filter='"+file_filter+"'; $d.Multiselect=$true; [Console]::OutputEncoding=[System.Text.Encoding]::UTF8; if($d.ShowDialog() -eq 'OK'){ ConvertTo-Json -Compress -InputObject @($d.FileNames) } else { Write-Output '[]' }"
                    proc=subprocess.run(['powershell','-NoProfile','-STA','-Command',script],capture_output=True,timeout=180,creationflags=getattr(subprocess,'CREATE_NO_WINDOW',0))
                    if proc.returncode:raise ValueError('附件选择窗口未能打开')
                    return self.reply(value={'paths':json.loads(proc.stdout.decode('utf-8-sig','replace').strip() or '[]')})
                if mode=='videos':
                    script="Add-Type -AssemblyName System.Windows.Forms; $d=New-Object System.Windows.Forms.OpenFileDialog; $d.Filter='Video files|*.mp4;*.mov;*.mkv;*.avi;*.m4v;*.webm;*.ts'; $d.Multiselect="+('$false' if body.get('multiple') is False else '$true')+"; [Console]::OutputEncoding=[System.Text.Encoding]::UTF8; if($d.ShowDialog() -eq 'OK'){ ConvertTo-Json -Compress -InputObject @($d.FileNames) } else { Write-Output '[]' }"
                    proc=subprocess.run(['powershell','-NoProfile','-STA','-Command',script],capture_output=True,timeout=180,creationflags=getattr(subprocess,'CREATE_NO_WINDOW',0))
                    if proc.returncode:raise ValueError('视频选择窗口未能打开，可以在文件名工具中粘贴本地路径')
                    return self.reply(value={'paths':json.loads(proc.stdout.decode('utf-8-sig','replace').strip() or '[]')})
                script="Add-Type -AssemblyName System.Windows.Forms; "
                if mode in ['template','python','image','engine','audio']:
                    script+="$d=New-Object System.Windows.Forms.OpenFileDialog; $d.Filter='"+('Audio|*.mp3;*.wav;*.m4a;*.aac;*.flac;*.ogg' if mode=='audio' else 'AI engine (*.zip)|*.zip' if mode=='engine' else 'Word template (*.docx)|*.docx' if mode=='template' else 'Images (*.jpg;*.jpeg;*.png;*.webp)|*.jpg;*.jpeg;*.png;*.webp' if mode=='image' else 'Python runtime (*.exe)|*.exe')+"';"
                else:script+="$d=New-Object System.Windows.Forms.FolderBrowserDialog; $d.Description='选择文件夹';"
                script+="if($d.ShowDialog() -eq 'OK'){"
                script+="[Console]::OutputEncoding=[System.Text.Encoding]::UTF8; Write-Output $d."+('FileName' if mode in ['template','python','image','engine','audio'] else 'SelectedPath')+'}'
                proc=subprocess.run(['powershell','-NoProfile','-STA','-Command',script],capture_output=True,timeout=180,
                    creationflags=getattr(subprocess,'CREATE_NO_WINDOW',0))
                return self.reply(value={'path':proc.stdout.decode('utf-8-sig','replace').strip()})
            if path=='/api/open-output':
                target=project_dir(get_project(body['id']),STATE['config']) if body.get('id') else Path(STATE['config']['output']).resolve()
                target.mkdir(parents=True,exist_ok=True)
                if os.name=='nt':os.startfile(str(target))
                return self.reply(value={'path':str(target)})
            return self.reply(404,{'error':'接口不存在'})
        except (BrokenPipeError,ConnectionResetError):pass
        except PermissionError as e:self.reply(403,{'error':str(e)})
        except ai.ProviderError as e:self.reply(400,{'error':ai.redact(e),'code':e.code,'upstreamStatus':e.status})
        except Exception as e:self.reply(400,{'error':ai.redact(e)})

if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--port',type=int,default=8765);parser.add_argument('--open',action='store_true');args=parser.parse_args()
    try:server=ThreadingHTTPServer(('127.0.0.1',args.port),Handler)
    except OSError:
        print(f'端口 {args.port} 已占用。若已有本程序运行，请打开 http://127.0.0.1:{args.port} ，或使用 --port 指定其他端口。');sys.exit(1)
    print(f'漫剧工坊 http://127.0.0.1:{args.port}',flush=True)
    if args.open:webbrowser.open(f'http://127.0.0.1:{args.port}')
    try:server.serve_forever()
    except KeyboardInterrupt:pass
    finally:
        shutdown_jobs()
        server.server_close()
